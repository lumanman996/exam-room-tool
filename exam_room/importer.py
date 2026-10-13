"""Excel/CSV 预检、列对应、模板。保留考号前导零和原始列。"""
import csv
import io
import json
import re
from datetime import datetime, date, time
import openpyxl
from .model import ENTITY_DEFAULTS, normalize_plan, new_plan, split_values

FIELDS={
 'students':{'name':'姓名','grade':'年级','class_name':'班级','student_no':'学号','number':'考号','score':'成绩','subjects':'参试科目'},
 'rooms':{'code':'考场编号','name':'考场名称','location':'位置','rows':'行数','cols':'桌列数','per_desk':'每桌人数','capacity':'容量','order':'座号顺序','direction':'座号方向','aisles':'过道后列','disabled':'停用座号','availability':'可用时段','invigilators':'监考人数'},
 'teachers':{'name':'姓名','subjects':'任教学科','classes':'任教班级','unavailable':'不可监考时段'},
 'exams':{'grade':'年级','subject':'科目','duration':'考试时长','date':'日期','start':'开始时间','end':'结束时间','locked':'锁定时间','special_seating':'单独排座','invigilators':'监考人数'},
}
ALIASES={'name':['学生姓名','教师姓名','名字'],'class_name':['班级(匀班)','班级（匀班）','新班级'],
         'number':['准考证号'],'student_no':['学籍号'],'score':['总分','总成绩','总分(匀班)'],'subject':['学科'],
         'cols':['列数'],'duration':['时长','时长（分钟）']}


def scalar(cell):
    v=cell.value
    if isinstance(v,(datetime,date,time)): return v.isoformat().split('T')[-1] if isinstance(v,time) else v.isoformat().split('T')[0]
    if isinstance(v,(int,float)) and re.fullmatch(r'0+',cell.number_format or ''):
        return str(int(v)).zfill(len(cell.number_format))
    return v if v is not None else ''


def preview(blob,filename,kind,sheet=None,header_row=None):
    if kind not in FIELDS: raise ValueError('不支持这种资料。')
    if len(blob)>30*1024*1024: raise ValueError('文件超过 30 MB，请拆分后导入。')
    if filename.lower().endswith('.csv'):
        for encoding in ('utf-8-sig','gb18030'):
            try: text=blob.decode(encoding); break
            except UnicodeDecodeError: continue
        else: raise ValueError('无法读取 CSV 编码，请另存为 UTF-8。')
        tables={'CSV':list(csv.reader(io.StringIO(text)))}
    elif filename.lower().endswith('.xlsx'):
        try:
            wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
            tables={ws.title:[[scalar(c) for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}
            wb.close()
        except Exception as e: raise ValueError('无法读取 Excel，请确认文件格式为 .xlsx。') from e
    else: raise ValueError('请选择 .xlsx 或 .csv 文件。')
    chosen=sheet if sheet in tables else next((s for s in ('总名单','成绩明细') if s in tables),next(iter(tables)))
    rows=tables[chosen]
    if not rows: raise ValueError('工作表没有数据。')
    labels=set(FIELDS[kind].values())|{v for vs in ALIASES.values() for v in vs}
    hr=int(header_row or 0)
    if not hr:
        hr=1+max(range(min(20,len(rows))),key=lambda i:len(labels&{str(c).strip() for c in rows[i]}))
    if not 1<=hr<=len(rows): raise ValueError('表头行超出了工作表。')
    headers=[]; used=set()
    for i,c in enumerate(rows[hr-1]):
        name=str(c).strip() or f'未命名列{i+1}'
        if name in used: name+=f'（第{i+1}列）'
        used.add(name); headers.append(name)
    data=[{h:(r[j] if j<len(r) else '') for j,h in enumerate(headers)} for r in rows[hr:] if any(str(c).strip() for c in r)]
    mapping={}
    for key,label in FIELDS[kind].items():
        candidates=([*ALIASES.get(key,[]),label] if key=='class_name' else [label,*ALIASES.get(key,[])])
        mapping[key]=next((h for a in candidates for h in headers if re.sub(r'\s','',h)==a),'')
    return dict(sheets=list(tables),sheet=chosen,header_row=hr,headers=headers,rows=data,mapping=mapping,fields=FIELDS[kind],count=len(data))


def intervals(value):
    if not value: return []
    if isinstance(value,list): return value
    s=str(value).strip()
    if s.startswith('['): return json.loads(s)
    out=[]
    for part in re.split(r'[;；\n]+',s):
        m=re.fullmatch(r'\s*(\d{4}-\d{2}-\d{2})\s+(\d{1,2}:\d{2})\s*[-—～~]\s*(\d{1,2}:\d{2})\s*',part)
        if not m: raise ValueError('时段请写成 2026-11-12 08:00-12:00，多段用分号分开。')
        out.append(dict(date=m[1],start=m[2].zfill(5),end=m[3].zfill(5)))
    return out


def import_rows(kind,rows,mapping):
    if kind not in FIELDS: raise ValueError('资料类型不正确。')
    out=[]
    for i,r in enumerate(rows):
        d={k:r.get(col,'') for k,col in mapping.items() if col}
        d={k:v for k,v in d.items() if v not in ('',None)}
        if kind=='students':
            d['original']=r
            if not d.get('grade'):
                cls=str(d.get('class_name',''))
                d['grade']=next((g+'年级' for g in ('七','八','九') if g in cls),'')
        for k in ('availability','unavailable'):
            if k in d: d[k]=intervals(d[k])
        for k in ('locked','special_seating'):
            if k in d: d[k]=str(d[k]).lower() in ('是','true','1','锁定')
        if 'order' in d: d['order']='snake' if str(d['order']) in ('蛇形','snake') else 'straight'
        if 'direction' in d: d['direction']='vertical' if str(d['direction']) in ('竖排','竖向','vertical') else 'horizontal'
        if 'invigilators' in d: d['invigilators']=int(d['invigilators'])
        for k in ('start','end'):
            if k in d: d[k]=str(d[k])[:5].zfill(5)
        out.append(d)
    p=new_plan(); p[kind]=out
    return normalize_plan(p)[kind]


def template(kind):
    if kind not in FIELDS: raise ValueError('资料类型不正确。')
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='填写资料'; ws.append(list(FIELDS[kind].values()))
    examples={
      'students':['学生甲','七年级','七1','001','00001',450,''],
      'rooms':['01','第01考场','教学楼一层',5,6,1,30,'顺排','横排','3','','',2],
      'teachers':['教师甲','语文','七1、七2','2026-11-12 08:00-12:00'],
      'exams':['七年级','语文',90,'2026-11-12','08:30','10:00','否','否',2]}
    ws.append(examples[kind]); ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    from openpyxl.styles import Font,PatternFill
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='334C83')
    for c in ws[2]: c.fill=PatternFill('solid',fgColor='FFF4CC')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=20
    note=wb.create_sheet('填写说明'); note.append(['第二行是虚构示例，使用前请删除。'])
    note.append(['考号和学号请用文本格式，保留前导零。'])
    note.append(['多项用顿号分隔；时段用 YYYY-MM-DD HH:MM-HH:MM，多段用分号分隔。'])
    note.append(['过道后列：填写过道前一列的桌列号，例如 3 表示第 3、4 桌列之间有过道。'])
    note.column_dimensions['A'].width=100
    out=io.BytesIO(); wb.save(out); return out.getvalue()


def preflight(plan,kind,rows,mapping,mode='append'):
    """不写入。用唯一学号/考号等强身份键匹配，绝不按学生姓名合并。"""
    from copy import deepcopy
    from collections import Counter
    from .validation import data_issues
    if kind not in FIELDS or mode not in ('append','replace','update'):raise ValueError('导入类型或方式不正确。')
    old=plan[kind];errors=[];warnings=[];result=[];seen=set();row_for={};added=updated=0
    def keys(row):
        if kind=='students':return [(k,str(row.get(k,''))) for k in ('student_no','number') if row.get(k)]
        if kind=='rooms':return [('code',str(row.get('code','')).lstrip('0'))] if row.get('code') else []
        if kind=='teachers':return [('name',str(row.get('name','')).strip())] if row.get('name') else []
        return [('exam',str(row.get('grade'))+'|'+str(row.get('subject')))] if row.get('grade') and row.get('subject') else []
    index={}
    for row in old:
        for key in keys(row):index.setdefault(key,[]).append(row)
    required={'students':['name','grade','class_name'],'rooms':['code','name'],'teachers':['name'],'exams':['grade','subject','duration']}[kind]
    for i,raw in enumerate(rows,1):
        try:
            prepared=import_rows(kind,[raw],mapping)[0]
            matches={r['id']:r for key in keys(prepared) for r in index.get(key,[])} if mode!='append' else {}
            if len(matches)>1:raise ValueError('身份键匹配多条资料，请先核对唯一学号、考号或内部编号。')
            existing=next(iter(matches.values()),None)
            if mode=='update' and existing is None:raise ValueError('没有找到可唯一更新的资料；学生须使用已有学号或考号，不能仅凭姓名更新。')
            if existing:
                if existing['id'] in seen:raise ValueError('本次导入多行匹配同一条已有资料。')
                provided={k for k,v in mapping.items() if v}
                merged={**deepcopy(existing),**{k:prepared[k] for k in provided if k in prepared},'id':existing['id']}
                if kind=='students':merged['original']={**existing.get('original',{}),**raw}
                if kind=='exams':
                    from .subjects import normalize_exam
                    merged=normalize_exam(merged)
                prepared=merged;updated+=1
            else:added+=1
            if kind=='exams' and mapping.get('duration') and prepared.get('start') and prepared.get('end'):
                from .subjects import duration
                value=raw.get(mapping['duration'])
                if value not in ('',None) and duration(value)!=prepared['duration']:
                    errors.append(dict(row=i,field='duration',code='duration_mismatch',message=f'第 {i} 行：起止时间与考试时长不一致，请修正分钟数或取消时长列对应，由起止时间计算。'))
            seen.add(prepared['id'])
            missing=[k for k in required if not str(prepared.get(k,'')).strip()]
            if missing:
                for field in missing:errors.append(dict(row=i,field=field,code='required',message=f'第 {i} 行：请填写{FIELDS[kind][field]}。'))
            result.append(prepared);row_for[prepared['id']]=i
        except (ValueError,TypeError,OverflowError) as e:errors.append(dict(row=i,field='',code='row_invalid',message=f'第 {i} 行：{e}'))
    prospective=deepcopy(plan)
    if mode=='replace':prospective[kind]=result
    elif mode=='update':
        replacement={r['id']:r for r in result};prospective[kind]=[replacement.get(r['id'],r) for r in old]
    else:prospective[kind]=old+result
    # Structural checks and duplicate identity checks, not missing later workflow data.
    for issue in data_issues(prospective):
        if issue.get('entity')==kind and issue['code'] not in ('grade_no_exam','subject_unknown'):
            targets=issue.get('room_ids') or [issue.get('entity_id')]
            for ident in targets:errors.append(dict(**issue,row=row_for.get(ident,0)))
    if kind=='students':
        for key in ('student_no','number'):
            values={}
            for r in prospective[kind]:
                value=r.get(key)
                if value:
                    if value in values:errors.append(dict(code='duplicate',field=key,row=row_for.get(r['id'],0),message=f'{FIELDS[kind][key]} {value} 重复，请核对。'))
                    values[value]=r['id']
        counts=Counter(r.get('name') for r in prospective[kind])
        for name,count in counts.items():
            if name and count>1:warnings.append(dict(code='same_name',message=f'姓名“{name}”出现 {count} 次，已按独立编号保留。'))
    return dict(plan=prospective,rows=result,errors=errors,warnings=warnings,added=added,updated=updated,removed=len({r['id'] for r in old}-{r['id'] for r in prospective[kind]}),mode=mode,revision=plan['revision'])
