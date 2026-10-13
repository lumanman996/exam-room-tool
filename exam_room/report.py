"""所有材料从同一方案生成。正式导出先经过统一检查。"""
from collections import defaultdict
from copy import deepcopy
from html import escape
import io
import openpyxl
from openpyxl.styles import Alignment,Font,PatternFill
from . import engine

SECTIONS={'schedule':'考试日程','students':'学生编排总表','rooms':'考场使用表','invigilation':'监考总表',
          'room_students':'按考场学生总表','student_details':'学生科目明细','workload':'监考工作量','classes':'分班通知单','personal':'学生个人考试安排','tickets':'准考证',
          'rosters':'考场名单','seating':'座次表','doors':'门贴','labels':'桌贴','teacher_personal':'教师个人安排','signins':'考场签到表'}


def all_rows(plan):
    rooms={r['id']:r for r in plan['rooms']}; teachers={t['id']:t for t in plan['teachers']}
    ds=engine.duties(plan); duty_for={(e,d['room_id']):d for d in ds for e in d['exam_ids']}
    rows=[]
    for e in sorted(plan['exams'],key=lambda x:(x['date'],x['start'],x['grade'],x['subject'])):
        seats=plan['seats'].get(e['id'],{}) if e.get('special_seating') else plan['seats'].get('base',{})
        for s in engine.participants(plan,e):
            a=seats.get(s['id'],{}); r=rooms.get(a.get('room_id'),{})
            d=duty_for.get((e['id'],a.get('room_id')),{}); inv=plan['invigilation'].get(d.get('id'),[])
            names='、'.join(teachers.get(t.get('teacher_id'),{}).get('name','') for t in inv)
            rows.append(dict(student_id=s['id'],exam_id=e['id'],room_id=a.get('room_id',''),
                日期=e['date'],时间=f"{e['start']}—{e['end']}",年级=s['grade'],班级=s['class_name'],姓名=s['name'],考号=s['number'],
                科目=e['subject'],考场=r.get('name','未安排'),考场编号=r.get('code',''),位置=r.get('location',''),座号=a.get('seat',''),监考教师=names))
    return rows


def materials(plan,scope=None):
    from .printing import filtered_rows,room_blocks
    rows=filtered_rows(plan,scope) if scope else all_rows(plan); ds=engine.duties(plan); rooms={x['id']:x for x in plan['rooms']}; teachers={x['id']:x for x in plan['teachers']}
    student_cols=['年级','班级','姓名','考号','日期','时间','科目','考场','位置','座号']
    def table(cols,data): return {'columns':cols,'rows':[[seat_text(r.get(k,'')) if k=='座号' else r.get(k,'') for k in cols] for r in data]}
    if scope:
        from .printing import scope_options
        filters=scope_options(plan,scope)
        ds=[d for d in ds if any(r['room_id']==d['room_id'] and r['exam_id'] in d['exam_ids'] for r in rows)]
    duties=[]; personal=[]
    for d in ds:
        r=rooms.get(d['room_id'],{}); inv=plan['invigilation'].get(d['id'],[])
        names='、'.join(teachers.get(t.get('teacher_id'),{}).get('name','') for t in inv)
        duties.append(dict(日期=d['date'],时间=f"{d['start']}—{d['end']}",考场=r.get('name',''),位置=r.get('location',''),
                           科目='、'.join(d['subjects']),人数=len(d['student_ids']),应到监考=d['required'],监考教师=names))
        for t in inv:
            if t.get('teacher_id') in teachers:
                personal.append(dict(教师=teachers[t['teacher_id']]['name'],**duties[-1]))
    schedule=[dict(年级=e['grade'],科目=e['subject'],日期=e['date'],开始=e['start'],结束=e['end'],时长=e['duration']) for e in plan['exams'] if not scope or ((not filters['exam_ids'] or e['id'] in filters['exam_ids']) and (not filters['grades'] or e['grade'] in filters['grades']))]
    loads=[dict(教师=t['name'],场次=len(own),监考分钟=sum(engine._event(d)[2]-engine._event(d)[1] for d in own)) for t in plan['teachers'] for own in [[d for d in ds if any(a.get('teacher_id')==t['id'] for a in plan['invigilation'].get(d['id'],[]))]]]
    out={k:table(student_cols,rows) for k in ('students','classes','personal','tickets','rosters','seating','labels')}
    out.update(schedule=table(['年级','科目','日期','开始','结束','时长'],schedule),
               rooms=table(['日期','时间','考场','位置','科目','人数'],duties),
               invigilation=table(['日期','时间','考场','科目','人数','应到监考','监考教师'],duties),
               workload=table(['教师','场次','监考分钟'],loads),
               doors=table(['日期','时间','考场','位置','科目','人数'],duties),
               teacher_personal=table(['教师','日期','时间','考场','位置','科目'],personal),
               signins=table(student_cols+['签到'],rows))
    blocks=room_blocks(plan,scope);out['room_students']=table(['考场编号','考场','位置','范围','座号','考号','姓名','班级','年级'],[{**r,'范围':b['scope']} for b in blocks for r in b['rows']])
    out['student_details']=deepcopy(out['students'])
    return out,rows,ds


def issue_text(plan,issue):
    labels=[]
    mappings={'student':('students',lambda x:x.get('class_name','')+' '+x.get('name','')+' '+x.get('number','')), 'teacher':('teachers',lambda x:x.get('name','')), 'room':('rooms',lambda x:'考场 '+x.get('code','')+' '+x.get('name','')), 'exam':('exams',lambda x:x.get('grade','')+' '+x.get('subject','')+' '+x.get('date','')+' '+x.get('start',''))}
    for kind,(entity,label) in mappings.items():
        ids=issue.get(kind+'_ids',[])+([issue[kind+'_id']] if issue.get(kind+'_id') else [])
        found={x['id']:x for x in plan[entity]}
        labels.extend(label(found[ident]).strip() for ident in ids if ident in found)
    if issue.get('seat'):labels.append('座号 '+seat_text(issue['seat']))
    if issue.get('duty_id'):
        d=next((d for d in engine.duties(plan) if d['id']==issue['duty_id']),None)
        if d:
            room=next((r for r in plan['rooms'] if r['id']==d['room_id']),{})
            labels.append('考场 '+room.get('code','')+' '+d['date']+' '+d['start']+'—'+d['end'])
    return issue['message']+('：'+'；'.join(dict.fromkeys(labels)) if labels else '')


def export_stages(key):
    if key=='schedule':return {'data','schedule'}
    if key in ('invigilation','workload','teacher_personal','signins'):return {'data','schedule','seats','numbers','monitors'}
    return {'data','schedule','seats','numbers'}


def check_export(plan,draft,sections=None):
    result=engine.validate(plan)
    if sections:
        allowed=set().union(*(export_stages(key) for key in sections))
        result['all_errors']=list(result['errors'])
        result['errors']=[e for e in result['errors'] if e['stage'] in allowed]

    if result['errors'] and not draft: raise ValueError(f"还有 {len(result['errors'])} 项问题需要处理。请先到「编排检查」，或选择草稿导出。")
    return result


def build_excel(plan,sections,draft=False,layouts=None,scope=None):
    from . import printing
    from openpyxl.worksheet.pagebreak import Break
    result=check_export(plan,draft,sections); mats,rows,ds=materials(plan,scope)
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    for key in sections:
        if key not in SECTIONS: raise ValueError('所选导出材料已不存在，请重新选择。')
        m=deepcopy(mats[key]);layout=printing.layout_for(plan,key,layouts)
        configured=bool(layouts or scope or plan.get('report_layouts') or key=='room_students')
        grouped=key in ('room_students','classes','personal','tickets','labels','rosters','signins','teacher_personal')
        if configured and key!='seating':
            data=[dict(zip(m['columns'],row)) for row in m['rows']]
            if grouped:
                blocks=printing.groups_for(plan,key,mats,rows,scope);available=blocks[0]['columns'] if blocks else printing.default_layouts()[key]['field_order'];data=[r for b in blocks for r in b['rows']]
            else:available=m['columns']
            fields,widths=printing.table_spec(layout,available,data);m=dict(columns=fields,rows=[[r.get(k,'') for k in fields] for r in data])
        ws=wb.create_sheet(SECTIONS[key]); ws.append([('【草稿】' if draft else '')+plan['name']+' · '+SECTIONS[key]])
        ws.append(m['columns'])
        for row in m['rows']:
            # Explicitly store imported/user strings as text, never Excel formulas.
            ws.append(row)
            for c in ws[ws.max_row]:
                if isinstance(c.value,str): c.data_type='s'
        ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{openpyxl.utils.get_column_letter(len(m["columns"]))}{ws.max_row}'
        ws.print_title_rows='1:2'; ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        for c in ws[2]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='334C83')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(36,max(14,max(len(str(c.value or '')) for c in col)*1.5))
        for row in ws:
            for c in row:
                if isinstance(c.value,str):c.data_type='s'
                c.alignment=Alignment(vertical='center',wrap_text=True)
        if configured:
            ws.page_setup.orientation=layout['orientation'];ws.page_setup.paperSize=ws.PAPERSIZE_A3 if layout['paper']=='A3' else ws.PAPERSIZE_A4
            from openpyxl.worksheet.page import PageMargins
            ws.page_margins=PageMargins(**{k:v/25.4 for k,v in layout['margins'].items()})
            for row in ws:
                for c in row:c.font=Font(name='宋体',size=layout['font_size'],bold=c.row<=2,color='000000')
                ws.row_dimensions[row[0].row].height=layout['font_size']*layout['line_height']+9
            for c in ws[2]:c.fill=PatternFill('solid',fgColor='FFFFFF' if layout['color']=='mono' else 'EDF0F7')
            if key!='seating':
                for i,k in enumerate(m['columns'],1):ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=widths[k]*3.78/7
            if not grouped and key!='seating':
                for n,r in enumerate(data,3):ws.row_dimensions[n].height=max(ws.row_dimensions[n].height or 0,printing.row_height(r,m['columns'],widths,layout)*72/25.4)
            if grouped:
                ws.delete_rows(3,ws.max_row);cols=len(m['columns'])
                # Separate printable room blocks; consistent data and selected
                # fields. Excel uses the requested columns side by side.
                column_count=layout['columns'];rowtop=3;rowbottom=3
                for lane in range(1,column_count):ws.column_dimensions[openpyxl.utils.get_column_letter(lane*(cols+1))].width=6*3.78/7
                for index,b in enumerate(blocks):
                    lane=index%column_count if layout['page_break']=='flow' else 0
                    if lane==0:
                        rowtop=rowbottom;rowbottom=rowtop
                        if index and layout['page_break']=='room':ws.row_breaks.append(Break(id=rowtop-1))
                    offset=lane*(cols+1)
                    ws.cell(rowtop,offset+1,b['title']);ws.merge_cells(start_row=rowtop,start_column=offset+1,end_row=rowtop,end_column=offset+cols)
                    ws.cell(rowtop+1,offset+1,b['scope']);ws.merge_cells(start_row=rowtop+1,start_column=offset+1,end_row=rowtop+1,end_column=offset+cols)
                    for j,k in enumerate(m['columns'],1):ws.cell(rowtop+2,offset+j,k);ws.column_dimensions[openpyxl.utils.get_column_letter(offset+j)].width=widths[k]*3.78/7
                    for n,r in enumerate(b['rows'],rowtop+3):
                        for j,k in enumerate(m['columns'],1):ws.cell(n,offset+j,r.get(k,''))
                        ws.row_dimensions[n].height=max(ws.row_dimensions[n].height or 0,printing.row_height(r,m['columns'],widths,layout)*72/25.4)
                    for n,k in ((rowtop,'title'),(rowtop+1,'scope')):ws.row_dimensions[n].height=max(ws.row_dimensions[n].height or 0,printing.row_height({k:b[k]},[k],{k:sum(widths.values())},layout)*72/25.4)
                    rowbottom=max(rowbottom,rowtop+len(b['rows'])+5)
                ws.auto_filter.ref=None
                for row in ws:
                    for c in row:
                        if isinstance(c.value,str):c.data_type='s'
                        c.font=Font(name='宋体',size=layout['font_size']);c.alignment=Alignment(wrap_text=True,vertical='center')
                    if not ws.row_dimensions[row[0].row].height:ws.row_dimensions[row[0].row].height=layout['font_size']*layout['line_height']+9
            last_column=(len(m['columns'])+1)*layout['columns']-1 if grouped else len(m['columns'])
            ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(1,last_column))
            _,_,page_width,_=printing.dimensions(layout)
            ws.row_dimensions[1].height=printing.row_height({'title':ws['A1'].value},['title'],{'title':page_width},layout)*72/25.4
            ws.print_options.horizontalCentered=True
        ws.oddFooter.center.text='考场编排 · 方案版本 '+str(plan['revision'])+' · 第 &P 页'
    if not wb.worksheets: raise ValueError('请至少选择一种导出材料。')
    ws=wb.create_sheet('编排检查');ws.append(['级别','说明'])
    for level,key in [('必须处理','errors'),('提示','warnings')]:
        for issue in result[key]: ws.append([level,issue_text(plan,issue)])
    for row in ws:
        for c in row:
            if isinstance(c.value,str): c.data_type='s'
    ws.column_dimensions['A'].width=15;ws.column_dimensions['B'].width=100
    out=io.BytesIO();wb.save(out);return out.getvalue()


def esc(x):return escape(str(x if x is not None else ''))

def seat_text(value):
    return str(value).zfill(2) if value not in ('',None) else ''

def html_table(cols,rows):
    return '<table><thead><tr>'+''.join('<th>'+esc(c)+'</th>' for c in cols)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(seat_text(v) if cols[i]=='座号' else v)+'</td>' for i,v in enumerate(r))+'</tr>' for r in rows)+'</tbody></table>'


def build_html(plan,sections,draft=False,landscape=True,layouts=None,scope=None):
    if layouts or scope or plan.get('report_layouts') or any(k in ('room_students','student_details') for k in sections):
        from .printing import document
        return document(plan,sections,draft,layouts,scope)['html']
    result=check_export(plan,draft,sections); mats,rows,ds=materials(plan); rooms={r['id']:r for r in plan['rooms']}
    body=[]
    def page(title,content,wide=False):
        body.append(f'<section class="page {"wide" if wide else ""}"><header><span>考场编排</span><span>{esc(plan["school"])}</span></header><h1>{esc(plan["name"])}</h1><h2>{esc(title)}</h2>{content}<footer>考场编排 · {"草稿，请核对后使用" if draft else "正式材料"}</footer></section>')
    cols=['班级','姓名','考号','科目','座号']
    for key in sections:
        if key not in SECTIONS:continue
        if key=='seating':
            for d in ds:
                room=rooms[d['room_id']]; actual=engine.seat_layout(room); bypos={(s['row'],s['col']):s for s in actual}
                physical=engine.seat_layout({**room,'disabled':[],'capacity':room['rows']*room['cols']*room['per_desk']})
                allpos={(s['row'],s['col']):s for s in physical}
                placed={}
                for r in rows:
                    if r['room_id']==d['room_id'] and r['exam_id'] in d['exam_ids']:placed[r['座号']]=r
                cells=''; nr=room['rows'];nc=room['cols']*room['per_desk']
                for rr in range(1,nr+1):
                    cells+='<tr>'
                    for cc in range(1,nc+1):
                        seat=bypos.get((rr,cc)); st=placed.get(seat['seat']) if seat else None
                        gap=cc%room['per_desk']==0 and cc//room['per_desk'] in room['aisles']
                        inactive=allpos.get((rr,cc),{}).get('seat')
                        state='停用' if inactive in room['disabled'] else '未使用'
                        cells+=f'<td class="seat {"aisle" if gap else ""}">'+(f'<small>{seat_text(seat["seat"])} 号</small><b>{esc(st["姓名"]) if st else "空位"}</b><span>{esc(st["班级"]) if st else ""}</span><span>{esc(st["考号"]) if st else ""}</span>' if seat else f'<small>{seat_text(inactive)} 号 · {state}</small>')+'</td>'
                    cells+='</tr>'
                page(f'{room["name"]} · {d["date"]} {d["start"]}—{d["end"]}',f'<div class="platform">讲台</div><table class="seats">{cells}</table>',landscape)
        elif key in ('tickets','labels'):
            grouped=defaultdict(list)
            for r in rows:grouped[r['student_id']].append(r)
            cards=[]
            for rs in grouped.values():
                s=rs[0]
                if key=='tickets':
                    content=html_table(['日期','时间','科目','考场','座号'],[[r[k] for k in ['日期','时间','科目','考场','座号']] for r in rs])
                    cards.append(f'<article class="ticket"><h3>{esc(s["姓名"])} · {esc(s["班级"])}</h3><p>考号：{esc(s["考号"])}</p>{content}</article>')
                else:
                    locs=defaultdict(list)
                    for r in rs:locs[(r['room_id'],r['座号'])].append(r['科目'])
                    for (rid,seat),subs in locs.items():
                        cards.append(f'<article class="label"><small>{esc(rooms.get(rid,{}).get("name","未安排"))} · {esc(seat_text(seat))} 号</small><h3>{esc(s["姓名"])}</h3><p>{esc(s["班级"])}　{esc(s["考号"])}</p><small>{esc("、".join(subs))}</small></article>')
            page(SECTIONS[key],'<div class="cards">'+''.join(cards)+'</div>')
        elif key in ('rosters','signins','doors'):
            for d in ds:
                room=rooms[d['room_id']]; rs=[r for r in rows if r['room_id']==d['room_id'] and r['exam_id'] in d['exam_ids']]
                rs.sort(key=lambda r:int(r['座号'] or 0)); c=cols+(['签到'] if key=='signins' else [])
                if key=='doors':
                    content=f'<div class="door"><h2>{esc(room["name"])}</h2><p>{esc(room["location"])}</p><h3>{esc("、".join(d["subjects"]))}</h3><p>{esc(d["date"])}　{esc(d["start"])}—{esc(d["end"])}</p><p>共 {len(d["student_ids"])} 人</p></div>'
                else: content=html_table(c,[[r.get(k,'') for k in c] for r in rs])
                page(f'{SECTIONS[key]} · {room["name"]} · {d["date"]} {d["start"]}',content)
        elif key in ('classes','personal'):
            grouped=defaultdict(list)
            for r in rows:grouped[(r['年级'],r['班级']) if key=='classes' else (r['student_id'],)].append(r)
            for rs in grouped.values():
                m=mats[key];title=rs[0]['班级'] if key=='classes' else rs[0]['班级']+' '+rs[0]['姓名']
                page(SECTIONS[key]+' · '+title,html_table(m['columns'],[[r.get(k,'') for k in m['columns']] for r in rs]))
        elif key=='teacher_personal':
            m=mats[key]; grouped=defaultdict(list)
            # Teacher names may coincide; derive groups from teacher IDs.
            for t in plan['teachers']:
                dd=[d for d in ds if any(x['teacher_id']==t['id'] for x in plan['invigilation'].get(d['id'],[]))]
                if dd:page('教师安排 · '+t['name'],html_table(['日期','时间','考场','位置','科目'],[[d['date'],d['start']+'—'+d['end'],rooms[d['room_id']]['name'],rooms[d['room_id']]['location'],'、'.join(d['subjects'])] for d in dd]))
        else:
            m=mats[key];page(SECTIONS[key],html_table(m['columns'],m['rows']))
    if not body: raise ValueError('请至少选择一种材料。')
    if draft and result['errors']:page('草稿检查清单',html_table(['问题'],[[issue_text(plan,x)] for x in result['errors']]))
    css='''@page{size:A4;margin:13mm}@page wide{size:A4 landscape;margin:13mm}*{box-sizing:border-box}body{font-family:"Songti SC","SimSun",serif;color:#192437;margin:0;font-size:10pt}.page{break-after:page;position:relative}.page:last-child{break-after:auto}.wide{page:wide}header,footer{display:flex;justify-content:space-between;color:#667085;font-size:9pt;padding:6px 0;border-bottom:1px solid #cdd3de}footer{border:0;margin-top:16px}h1{font-size:20pt;margin:14px 0 6px}h2{font-size:13pt;margin:4px 0 16px}h3{margin:7px 0}table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{border:1px solid #ccd1dc;padding:5px;overflow-wrap:anywhere;text-align:center}th{background:#edf0f7}thead{display:table-header-group}tr{break-inside:avoid}.platform{text-align:center;border:1px solid #8796b4;padding:8px;margin:12px 30%}.seat{height:60px}.seat b,.seat span,.seat small{display:block}.seat small{color:#6b7280}.aisle{border-right:7px solid #eee}.cards{display:grid;grid-template-columns:1fr 1fr;gap:12px}.ticket,.label{border:1px solid #9daac2;padding:12px;break-inside:avoid}.label{text-align:center;min-height:115px}.ticket table{font-size:8pt}.door{text-align:center;padding:50px 0}.door h2{font-size:42pt}.door h3{font-size:25pt}.door p{font-size:18pt}.watermark{position:fixed;top:40%;left:15%;transform:rotate(-25deg);font-size:100pt;color:rgba(180,30,30,.09);z-index:-1}@media screen{body{background:#e8e9ec}.page{background:white;max-width:1100px;margin:20px auto;padding:35px;box-shadow:0 1px 10px #ccc}}'''
    return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>'+esc(plan['name'])+'</title><style>'+css+'</style><body>'+('<div class="watermark">草稿</div>' if draft else '')+''.join(body)+'</body></html>'
