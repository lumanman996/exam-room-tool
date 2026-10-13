from copy import deepcopy
import io
import pytest
import openpyxl
from exam_room import report,engine
from exam_room.model import normalize_plan
from test_backend_hardening import ready


def test_fixed_room_total_deduplicates_subject_rows_and_details_stay_separate():
    p=ready();p['exams'].append(dict(p['exams'][0],id='e2',date='2026-11-02',start='08:00',end='09:30'))
    p['settings']['slots']=[]
    mats,rows,ds=report.materials(p)
    assert len(mats['room_students']['rows'])==6
    assert len(mats['student_details']['rows'])==12


def test_layout_parameters_change_html_excel_and_persist(tmp_path):
    from exam_room.server import App
    from exam_room.licensing import License
    p=ready();layout={'paper':'A3','orientation':'portrait','columns':1,'font_size':14,'line_height':1.6,'margins':{'top':18,'right':16,'bottom':18,'left':16},'color':'mono','page_break':'room','field_order':['座号','姓名','考号','班级'],'column_widths':{'考号':55}}
    app=App(tmp_path,License(tmp_path/'lic',None));app.commit(p);rev=app.plan['revision']
    app.api_report_layout_save({'layouts':{'room_students':layout},'revision':rev})
    assert not app.plan.get('stale')
    loaded=app.store.load(p['id']);assert loaded['report_layouts']['room_students']['font_size']==14
    html=report.build_html(loaded,['room_students'])
    assert 'A3 portrait' in html and '14pt' in html and '55mm' in html and '18mm 16mm 18mm 16mm' in html
    wb=openpyxl.load_workbook(io.BytesIO(report.build_excel(loaded,['room_students'])))
    ws=wb[report.SECTIONS['room_students']]
    assert str(ws.page_setup.paperSize)==ws.PAPERSIZE_A3 and ws.page_setup.orientation=='portrait'
    assert ws.page_margins.top==pytest.approx(18/25.4)
    assert any(c.font.sz==14 for row in ws for c in row if c.value=='学生甲' or (isinstance(c.value,str) and '测试学生' in c.value))


def test_invalid_layout_does_not_silently_shrink_or_drop_students():
    p=ready()
    with pytest.raises(ValueError,match='栏|列宽|宽度'):
        report.build_html(p,['room_students'],layouts={'room_students':{'paper':'A4','orientation':'portrait','columns':3,'column_widths':{'考号':150}}})
    with pytest.raises(ValueError):report.build_html(p,['room_students'],layouts={'room_students':{'font_size':5}})


def test_reused_fixed_seat_gets_distinct_time_blocks_and_special_scope_uses_effective_seats():
    p=ready();p['settings']['slots']=[];p['students'].append(dict(p['students'][0],id='s8',name='虚构八甲',grade='八年级',class_name='八1',number='800001'))
    p['exams'].append(dict(p['exams'][0],id='eight',grade='八年级',start='10:00',end='11:30'))
    p['seats']['base']['s8']=deepcopy(p['seats']['base'][p['students'][0]['id']])
    from exam_room.printing import room_blocks
    blocks=room_blocks(p)
    assert len(blocks)==2 and all(b['scope'] for b in blocks)
    assert sum(len(b['rows']) for b in blocks)==7
    sid=p['students'][0]['id'];ex=p['exams'][0];old=p['seats']['base'][sid]['seat'];ex['special_seating']=True
    p['seats'][ex['id']]=deepcopy(p['seats']['base']);other=p['students'][1]['id'];p['seats'][ex['id']][sid],p['seats'][ex['id']][other]=p['seats'][ex['id']][other],p['seats'][ex['id']][sid]
    scoped=room_blocks(p,{'mode':'exam','exam_ids':[ex['id']]})
    row=next(r for b in scoped for r in b['rows'] if r['student_id']==sid)
    assert row['座号']!=str(old).zfill(2)


def test_async_export_pins_revision_does_not_block_edits_and_cancels(tmp_path,monkeypatch):
    import threading,time
    from exam_room.server import App
    from exam_room.licensing import License
    from exam_room import exporting
    app=App(tmp_path,License(tmp_path/'lic',None));app.commit(ready());old=deepcopy(app.plan)
    entered=threading.Event();release=threading.Event();original=exporting.build
    def slow(*args):entered.set();assert release.wait(5);return original(*args)
    monkeypatch.setattr(exporting,'build',slow)
    job=app.api_export_start({'format':'preview','sections':['room_students'],'revision':old['revision']})
    assert entered.wait(2)
    modified=deepcopy(app.plan);modified['name']='虚构修改后标题';app.api_save({'plan':modified,'revision':modified['revision']})
    release.set()
    for _ in range(100):
        state=app.api_export_status({'id':job['id']})
        if state['state']!='running':break
        time.sleep(.02)
    assert state['state']=='done',state
    assert state['result']['stale'] and state['result']['snapshot_revision']==old['revision']
    assert old['name'] in state['result']['html'] and modified['name'] not in state['result']['html']
    assert state['result']['page_count']==len(state['result']['pages'])
    entered.clear();release.clear();downloads=len(app.downloads)
    job=app.api_export_start({'format':'xlsx','sections':['students']});assert entered.wait(2)
    app.api_export_cancel({'id':job['id']});release.set()
    for _ in range(100):
        state=app.api_export_status({'id':job['id']})
        if state['state']!='running':break
        time.sleep(.02)
    assert state['state']=='cancelled' and len(app.downloads)==downloads


def test_related_material_gate_and_old_revision_are_checked_in_service(tmp_path):
    from exam_room.server import App
    from exam_room.licensing import License
    app=App(tmp_path,License(tmp_path/'lic',None));p=ready();p['invigilation']={};app.commit(p)
    assert app.api_export({'format':'xlsx','sections':['schedule']})['url']
    with pytest.raises(ValueError):app.api_export_start({'format':'preview','sections':['invigilation']})
    with pytest.raises(ValueError,match='更新'):app.api_export_start({'sections':['schedule'],'revision':-1})


def test_every_material_default_can_generate_html_and_excel_with_same_fields():
    from exam_room import printing
    p=ready();defaults=printing.default_layouts()
    for key in report.SECTIONS:
        doc=printing.document(p,[key],layouts={key:defaults[key]})
        assert doc['page_count']==len(doc['pages']) and 'section class="page' in doc['html']
        wb=openpyxl.load_workbook(io.BytesIO(report.build_excel(p,[key],layouts={key:defaults[key]})))
        assert report.SECTIONS[key] in wb.sheetnames


def test_excel_long_names_get_multiple_line_height_and_columns_have_six_mm_gap():
    p=ready();p['students'][0]['name']='虚构长姓名甲乙丙丁戊己庚辛壬癸子丑寅卯辰巳午未申酉戌亥'
    blob=report.build_excel(p,['room_students'],layouts={'room_students':{'columns':2}})
    ws=openpyxl.load_workbook(io.BytesIO(blob))[report.SECTIONS['room_students']]
    cell=next(c for row in ws for c in row if c.value==p['students'][0]['name'])
    assert ws.row_dimensions[cell.row].height>40
    assert ws.column_dimensions['F'].width==pytest.approx(6*3.78/7)
