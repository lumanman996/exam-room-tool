"""独立物理坐标答案核对实际座次图；不使用产品布局作为断言答案。"""
from copy import deepcopy
from html.parser import HTMLParser
import io
import re
import openpyxl
import pytest
from exam_room import printing,report
from exam_room.model import new_plan,normalize_plan


class SeatTables(HTMLParser):
    def __init__(self,html):
        super().__init__();self.tables=[];self.active=False;self.row=None;self.cell=None;self.depth=0;self.feed(html)
    def handle_starttag(self,tag,attrs):
        attrs=dict(attrs)
        if tag=='table' and 'seats' in attrs.get('class','').split():self.active=True;self.tables.append([])
        if not self.active:return
        if tag=='tr':self.row=[];self.tables[-1].append(self.row)
        if tag=='td':self.cell={'text':'','class':attrs.get('class','')};self.row.append(self.cell)
    def handle_endtag(self,tag):
        if tag=='td':self.cell=None
        if tag=='table':self.active=False
    def handle_data(self,data):
        if self.active and self.cell is not None:self.cell['text']+=data


def sample(direction='horizontal',order='straight',custom=True,special=False):
    # 2 physical rows, 3 desks per row, 2 seats per desk. This oracle is derived
    # from coordinates, never from seat_layout(), product state.layouts or HTML.
    labels={}
    for row in (1,2):
        for col in range(1,7):
            if direction=='horizontal':label=(row-1)*6+(col if order=='straight' or row%2 else 7-col)
            else:label=(col-1)*2+(row if order=='straight' or col%2 else 3-row)
            labels[row,col]=label
    if custom:labels[1,1],labels[1,3]=labels[1,3],labels[1,1]
    disabled_coord=(1,2);disabled=labels[disabled_coord]
    route=[]
    for line in range(1,3 if direction=='horizontal' else 7):
        offsets=range(1,7 if direction=='horizontal' else 3)
        for offset in offsets:
            step=offset if order=='straight' or line%2 else (7 if direction=='horizontal' else 3)-offset
            route.append((line,step) if direction=='horizontal' else (step,line))
    usable=set([pos for pos in route if pos!=disabled_coord][:10])
    p=new_plan('虚构座号打印回归');p['grades']=['七年级']
    p['rooms']=[dict(id='room',code='01',name='虚构双人桌考场',rows=2,cols=3,per_desk=2,capacity=10,disabled=[disabled],aisles=[1],direction=direction,order=order,
        seat_labels={f'{r}:{c}':n for (r,c),n in labels.items()} if custom else {})]
    p['students']=[dict(id=f's{i}',name=name,grade='七年级',class_name='七1',number=f'000{i+1}') for i,name in enumerate(['虚构甲','虚构乙','虚构丙'])]
    p['exams']=[dict(id='exam',grade='七年级',subject='虚构科目',duration=60,date='2026-11-17',start='08:00',end='09:00',special_seating=special)]
    fixed=[(1,1),(1,3),(2,2)];active_positions=[(2,2),(1,1),(1,3)] if special else fixed
    p['seats']['base']={f's{i}':dict(room_id='room',seat=labels[pos],locked=i==0) for i,pos in enumerate(fixed)}
    if special:p['seats']['exam']={f's{i}':dict(room_id='room',seat=labels[pos],locked=i==0) for i,pos in enumerate(active_positions)}
    occupants={pos:p['students'][i] for i,pos in enumerate(active_positions)}
    return normalize_plan(p),dict(labels=labels,usable=usable,disabled=disabled_coord,occupants=occupants)


def assert_physical_html(html,oracle):
    tables=SeatTables(html).tables;assert len(tables)==1
    rows=tables[0];assert len(rows)==2
    for row,cells in enumerate(rows,1):
        assert len(cells)==6
        for col,cell in enumerate(cells,1):
            coordinate=(row,col);text=re.sub(r'\s+','',cell['text'])
            assert text.startswith(f'{oracle["labels"][coordinate]:02}号'),f'{coordinate}: {text}'
            student=oracle['occupants'].get(coordinate)
            if student:assert student['name'] in text and student['number'] in text
            else:assert ('停用' if coordinate==oracle['disabled'] else '空位' if coordinate in oracle['usable'] else '未使用') in text
            assert ('aisle' in cell['class'].split())==(col==2)


@pytest.mark.parametrize('direction',['horizontal','vertical'])
@pytest.mark.parametrize('order',['straight','snake'])
@pytest.mark.parametrize('custom',[False,True])
@pytest.mark.parametrize('special',[False,True])
def test_actual_html_columns_are_physical_for_every_route_and_subject(direction,order,custom,special):
    p,oracle=sample(direction,order,custom,special);before=deepcopy(p)
    actual=printing.document(p,['seating'],True,scope={'mode':'exam','exam_ids':['exam']})
    assert_physical_html(actual['html'],oracle)
    # The legacy printable HTML branch has its own physical-coordinate traversal.
    assert_physical_html(report.build_html(p,['seating'],True),oracle)
    assert p==before and p['seats']['base']['s0']['locked']
    if special:assert p['seats']['exam']['s0']['locked']


@pytest.mark.parametrize('special',[False,True])
def test_actual_excel_seating_table_keeps_identity_labels_and_coordinates(special):
    p,oracle=sample(custom=True,special=special);before=deepcopy(p)
    ws=openpyxl.load_workbook(io.BytesIO(report.build_excel(p,['seating'],True,scope={'mode':'exam','exam_ids':['exam']})))[report.SECTIONS['seating']]
    headers=[c.value for c in ws[2]];rows=[dict(zip(headers,[c.value for c in row])) for row in list(ws.rows)[2:]]
    # Current Excel seating output is a filterable detail table, not a grid.
    assert len(rows)==3 and '座号' in headers
    for coordinate,student in oracle['occupants'].items():
        r=next(r for r in rows if r['考号']==student['number'])
        assert r['姓名']==student['name'] and r['座号']==f'{oracle["labels"][coordinate]:02}'
    assert p==before
