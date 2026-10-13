import io
import openpyxl
import pytest
from exam_room import engine,report
from exam_room.model import demo_plan

@pytest.fixture
def ready():
    p=demo_plan();p=engine.assign_seats(p)['plan'];p=engine.assign_invigilators(p)['plan'];return p


def test_all_materials_use_effective_subject_seats_and_hide_scores(ready):
    p=ready;e=p['exams'][0];students=engine.participants(p,e);a=p['seats']['base'][students[0]['id']];other=next(s for s in students[1:] if p['seats']['base'][s['id']]['room_id']==a['room_id'])
    pos=p['seats']['base'][other['id']]
    p=engine.adjust_seat(p,students[0]['id'],pos['room_id'],pos['seat'],exam_id=e['id'],swap=True)
    rows=report.all_rows(p)
    effective=next(r for r in rows if r['student_id']==students[0]['id'] and r['exam_id']==e['id'])
    assert effective['座号']==pos['seat']
    wb=openpyxl.load_workbook(io.BytesIO(report.build_excel(p,list(report.SECTIONS))))
    assert set(report.SECTIONS.values())<=set(wb.sheetnames)
    for sheet in wb:
        assert '成绩' not in [c.value for c in sheet[2]]
    html=report.build_html(p,list(report.SECTIONS))
    assert '讲台' in html and '准考证' in html and '成绩' not in html


def test_formal_export_rejects_missing_seats_draft_is_marked():
    p=demo_plan()
    with pytest.raises(ValueError):report.build_excel(p,['students'])
    w=openpyxl.load_workbook(io.BytesIO(report.build_excel(p,['students'],draft=True)))
    assert '草稿' in w['学生编排总表']['A1'].value
    assert '草稿' in report.build_html(p,['tickets'],draft=True)


def test_user_values_never_become_formulas_or_html(ready):
    ready['students'][0]['name']='=1+1'
    ready['name']='=1+1'
    ready['school']='<script>alert(1)</script>'
    wb=openpyxl.load_workbook(io.BytesIO(report.build_excel(ready,['students'])),data_only=False)
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)
    html=report.build_html(ready,['tickets'])
    assert '<script>alert' not in html and '&lt;script&gt;' in html


def test_two_digit_seat_numbers_and_unused_seats_are_explicit(ready):
    rows=report.all_rows(ready);expected=next(r for r in rows if r['座号']<10)
    mats,_,_=report.materials(ready)
    i=mats['students']['columns'].index('座号')
    assert all(isinstance(row[i],str) and len(row[i])==2 for row in mats['students']['rows'])
    html=report.build_html(ready,['seating','labels'])
    assert f'{expected["座号"]:02} 号' in html
    room=ready['rooms'][0];room['capacity']=29
    draft=report.build_html(ready,['seating'],draft=True)
    assert '未使用' in draft
