"""验测台自身不能把漏排、教师冲突、错误报表或挂起误报为通过。"""
from copy import deepcopy
import json
from pathlib import Path
import sys
import hashlib

import pytest

ROOT = Path(__file__).resolve().parents[1]


def test_fingerprint_uses_posix_keys_even_for_windows_relative_paths(tmp_path,monkeypatch):
    from pathlib import PureWindowsPath
    from tools import run_testbench
    p=tmp_path/'exam_room'/'虚构.py';p.parent.mkdir();p.write_bytes(b'fictional fingerprint')
    original=Path.relative_to
    monkeypatch.setattr(run_testbench,'ROOT',tmp_path)
    monkeypatch.setattr(Path,'relative_to',lambda self,root:PureWindowsPath(*original(self,root).parts))
    assert run_testbench.source_fingerprint()=={'exam_room/虚构.py':hashlib.sha256(p.read_bytes()).hexdigest()}


def test_windows_autocrlf_checkout_keeps_declared_fixture_bytes(tmp_path):
    import shutil,subprocess
    source=tmp_path/'source';source.mkdir()
    shutil.copyfile(ROOT/'.gitattributes',source/'.gitattributes')
    names=['testkit/fixtures/C01_tiny_completed/golden_expected.json',
           'testkit/fixtures/C01_tiny_completed/学生名单.csv',
           'testkit/fixtures/C01_tiny_completed/学生名单.xlsx']
    for name in names:
        target=source/name;target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copyfile(ROOT/name,target)
    def git(*args):
        return subprocess.run(['git','-C',str(source),'-c','core.autocrlf=true',*args],
            capture_output=True,check=True,timeout=15)
    git('init','--quiet');git('add','.gitattributes','testkit')
    checkout=tmp_path/'checkout'
    git('checkout-index','--all','--prefix='+checkout.as_posix()+'/')
    for name in names:assert (checkout/name).read_bytes()==(ROOT/name).read_bytes(),name


@pytest.mark.parametrize('status',['partial','feasible'])
def test_larger_scale_budget_still_rejects_any_unfilled_monitor(tmp_path,monkeypatch,status):
    from contextlib import contextmanager
    from types import SimpleNamespace
    from tools.testbench import scenarios
    calls=[]
    class Client:
        def solve(self,kind,limit):
            calls.append((kind,limit))
            return {'task':{'status':'feasible','unfilled':[]} if kind=='seats' else
                {'status':status,'unfilled':[{'missing':5}],'proved_minimum':False,'elapsed':90}}
        def call(self,*args):raise AssertionError('缺岗不能进入确认或导出步骤')
    @contextmanager
    def fake_service(*args):yield Client()
    monkeypatch.setattr(scenarios,'service',fake_service)
    with pytest.raises(AssertionError,match='unfilled'):
        scenarios.scale_case(SimpleNamespace(root=tmp_path),'SCALE-mixed')
    assert calls==[('seats',30),('invigilators',90)]
    terminal=json.loads((tmp_path/'SCALE-mixed/监考任务终态.json').read_text(encoding='utf-8'))
    assert terminal['unfilled']==[{'missing':5}]


def ready():
    return json.loads((ROOT/'testkit/fixtures/C01_tiny_completed/方案.json').read_text())


def test_independent_check_accepts_golden_and_rejects_duplicated_seat():
    from tools.testbench.common import assert_plan
    p = ready(); assert_plan(p)
    p['seats']['base']['s7-002']['seat'] = 1
    with pytest.raises(AssertionError, match='重复占座'):
        assert_plan(p)


def test_independent_check_rejects_missing_student_and_teacher_overlap():
    from tools.testbench.common import assert_plan
    p = ready(); p['seats']['base'].pop('s7-001')
    with pytest.raises(AssertionError, match='漏排'):
        assert_plan(p)
    p = ready(); p['invigilation']['r02|2026-11-12|08:30|10:00'][0]['teacher_id'] = 't01'
    with pytest.raises(AssertionError, match='教师时间冲突'):
        assert_plan(p)


def test_independent_check_rejects_subject_avoidance():
    from tools.testbench.common import assert_plan
    p = ready(); p['teachers'][0]['subjects'] = ['语文']
    with pytest.raises(AssertionError, match='任教学科'):
        assert_plan(p)


def test_excel_check_rejects_missing_row_and_changed_number():
    import io
    import openpyxl
    from exam_room.report import build_excel
    from tools.testbench.common import assert_excel
    p = ready(); blob = build_excel(p, ['students'])
    assert_excel(blob, p, ['students'])
    wb = openpyxl.load_workbook(io.BytesIO(blob)); ws = wb['学生编排总表']
    ws.delete_rows(3); out = io.BytesIO(); wb.save(out)
    with pytest.raises(AssertionError, match='明细'):
        assert_excel(out.getvalue(), p, ['students'])
    wb = openpyxl.load_workbook(io.BytesIO(blob)); ws = wb['学生编排总表']
    ws.cell(3, 4, '9999999'); out = io.BytesIO(); wb.save(out)
    with pytest.raises(AssertionError, match='明细'):
        assert_excel(out.getvalue(), p, ['students'])


def test_report_escapes_html_and_does_not_pass_unrun(tmp_path):
    from tools.run_testbench import write_report, exit_status
    data = dict(mode='full', environment={}, cases=[dict(id='X', name='<script>x</script>',
        status='未运行', seconds=0, details={})], limitations=[])
    write_report(tmp_path, data)
    html = (tmp_path/'report.html').read_text()
    assert '<script>x</script>' not in html and '&lt;script&gt;' in html
    assert exit_status(data['cases']) != 0


def test_all_materials_checked_and_corrupt_workload_and_room_detected():
    import io
    import openpyxl
    from exam_room.report import build_excel, SECTIONS
    from tools.testbench.common import assert_excel
    p=ready(); sections=list(SECTIONS); blob=build_excel(p,sections)
    assert_excel(blob,p,sections)
    for sheet,cell,value in [('监考工作量','B3',99),('按考场学生总表','C6','虚构错名')]:
        wb=openpyxl.load_workbook(io.BytesIO(blob)); wb[sheet][cell]=value
        out=io.BytesIO(); wb.save(out)
        with pytest.raises(AssertionError,match='明细'):
            assert_excel(out.getvalue(),p,sections)


def test_grouped_excel_checks_card_identity_and_missing_student():
    import io
    import openpyxl
    from exam_room.report import build_excel
    from tools.testbench.common import assert_excel
    p=ready(); sections=['classes','tickets','teacher_personal','signins']
    blob=build_excel(p,sections,scope={'mode':'fixed'})
    assert_excel(blob,p,sections)
    wb=openpyxl.load_workbook(io.BytesIO(blob));wb['准考证']['A3']='虚构错误学生 · 错班 · 考号 999'
    out=io.BytesIO();wb.save(out)
    with pytest.raises(AssertionError,match='分组明细'):
        assert_excel(out.getvalue(),p,sections)


def test_pdf_text_normalizes_font_radicals_and_line_breaks_but_rejects_missing_name():
    from tools.testbench.common import assert_pdf_text
    p={'students':[{'name':'虚构学生七001','number':'0010101'}]}
    assert_pdf_text('虚构学⽣七\n001 0010101',p)
    with pytest.raises(AssertionError,match='姓名'):
        assert_pdf_text('虚构学⽣七\n002 0010101',p)


def test_scoped_excel_uses_special_seats_and_rejects_wrong_time_caption():
    import io,openpyxl
    from exam_room.report import build_excel
    from tools.testbench.common import assert_excel,load
    p=load('C05_special_subject');e=next(e for e in p['exams'] if e['special_seating'])
    scope={'mode':'exam','exam_ids':[e['id']]}
    sections=['room_students','student_details'];blob=build_excel(p,sections,scope=scope)
    assert_excel(blob,p,sections,scope=scope)
    wb=openpyxl.load_workbook(io.BytesIO(blob));wb['按考场学生总表']['A4']='2099-01-01 08:00—09:00'
    out=io.BytesIO();wb.save(out)
    with pytest.raises(AssertionError,match='明细'):
        assert_excel(out.getvalue(),p,sections,scope=scope)


def test_subprocess_timeout_is_recorded_as_failure(tmp_path):
    from tools.run_testbench import command_result
    result = command_result([sys.executable, '-c', 'import time; time.sleep(30)'], tmp_path, 'hang', .15)
    assert result['status'] == '失败'
    assert result['details']['timeout']


def test_missing_command_is_blocked_with_log_and_not_passed(tmp_path):
    from tools.run_testbench import command_result,exit_status
    result=command_result([str(tmp_path/'fictional-missing-executable')],tmp_path,'missing',1)
    assert result['status']=='环境阻塞' and result['details']['exit_code'] is None
    assert Path(result['details']['log']).read_text() and exit_status([result])==1


def test_unknown_selected_case_cannot_pass_using_only_source_fingerprint(tmp_path):
    import subprocess
    result=subprocess.run([sys.executable,str(ROOT/'tools/run_testbench.py'),'--phase','scenarios',
        '--output',str(tmp_path),'--case','FICTIONAL-NO-SUCH-CASE'],cwd=ROOT,capture_output=True,text=True,timeout=15)
    assert result.returncode!=0,'一个场景都没执行，不能因为源码指纹一致而通过'
    data=json.loads((tmp_path/'report.json').read_text())
    assert data['summary']['失败']>0


def test_case_recorder_continues_after_failure_and_keeps_evidence(tmp_path):
    from tools.testbench.common import Recorder
    r = Recorder(tmp_path)
    def broken(): raise AssertionError('虚构失败')
    r.run('bad', '坏例', broken)
    r.run('good', '正常例', lambda: {'checked': True})
    assert [x['status'] for x in r.cases] == ['失败', '通过']
    assert (tmp_path/'bad/error.txt').is_file()
    assert '虚构失败' in (tmp_path/'bad/error.txt').read_text()


def test_proven_separate_grade_room_shortage_is_explained_before_solver(monkeypatch):
    from exam_room import engine
    from exam_room.model import demo_plan
    p=demo_plan(large=True); p['settings']['mix_grades']=False
    def unexpected_solver(*args, **kwargs):
        raise AssertionError('已能计算出至少51场，不应进入求解器等待超时')
    monkeypatch.setattr(engine,'_solve',unexpected_solver)
    result=engine.assign_seats(p)
    assert result['status']=='infeasible'
    assert '51' in result['message'] and '50' in result['message']


def test_browser_trace_scrubs_temporary_tokens_and_preserves_images(tmp_path):
    import zipfile
    from tools.testbench.common import scrub_trace
    trace=tmp_path/'trace.zip'
    with zipfile.ZipFile(trace,'w') as z:
        z.writestr('trace.network',json.dumps({'X-Token':'fictional-token'}))
        z.writestr('resources/image.png',b'\x89PNG\xff\x00')
    scrub_trace(trace,['fictional-token'])
    with zipfile.ZipFile(trace) as z:
        assert b'fictional-token' not in z.read('trace.network')
        assert z.read('resources/image.png')==b'\x89PNG\xff\x00'


def test_changed_source_fingerprint_cannot_be_reported_as_pass():
    from tools.run_testbench import source_stability
    c=source_stability({'app.py':'before'}, {'app.py':'after'})
    assert c['status']=='失败' and c['details']['changed_files']==['app.py']
