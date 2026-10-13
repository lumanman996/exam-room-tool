"""生成并检查虚构验测资料；不启动用户软件、不写用户数据、不执行最终验收。"""
from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import hashlib
import io
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill
from exam_room.model import demo_plan, new_plan, normalize_plan
from exam_room.importer import FIELDS

STAMP = '2026-10-06T00:00:00'
KINDS = ('students', 'rooms', 'teachers', 'exams')
LABELS = {'students': '学生名单', 'rooms': '考场资料', 'teachers': '教师资料', 'exams': '考试日程'}


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def fixed(plan, ident, name):
    p = normalize_plan(plan)
    p.update(id=ident, name=name, school='验测中学（完全虚构）',
             created_at=STAMP, updated_at=STAMP, revision=0)
    return p


def tiny():
    """手工给定位置和监考；不调用求解器或产品的 duties()。"""
    p = new_plan()
    for gi, grade in enumerate(('七', '八')):
        rid = f'r{gi + 1:02}'
        p['rooms'].append(dict(id=rid, code=f'{gi + 1:02}', name=f'虚构第{gi + 1:02}考场',
                               location='虚构教学楼一层', rows=2, cols=3, capacity=6))
        for i in range(6):
            sid = f's{gi + 7}-{i + 1:03}'
            p['students'].append(dict(id=sid, name=f'虚构学生{grade}{i + 1:03}', grade=grade+'年级',
                class_name=f'{grade}{i % 2 + 1}', student_no=f'00{gi + 7}{i + 1:03}',
                number=f'001{gi + 1:02}{i + 1:02}', score=400 + i * 10, subjects=[]))
            p['seats']['base'][sid] = dict(room_id=rid, seat=i+1, locked=False)
            p['number_sources'][sid] = dict(prefix='001', room_id=rid, number=f'001{gi + 1:02}{i + 1:02}')
        for j, subject in enumerate(('语文', '数学', '英语')):
            date = '2026-11-12' if j < 2 else '2026-11-13'
            start, end = ('08:30', '10:00') if j != 1 else ('14:00', '15:30')
            p['exams'].append(dict(id=f'e{gi + 7}-{j+1}', grade=grade+'年级', subject=subject,
                duration=90, date=date, start=start, end=end))
            p['invigilation'][f'{rid}|{date}|{start}|{end}'] = [
                dict(teacher_id=f't{gi*2+k+1:02}', locked=False) for k in range(2)]
    for i in range(4):
        p['teachers'].append(dict(id=f't{i+1:02}', name=f'虚构教师{i+1:03}', subjects=['体育'], classes=[]))
    p['settings'].update(number_prefix='001', number_rule='seat', seed=42,
        slots=[dict(date=d, start=a, end=b) for d in ('2026-11-12', '2026-11-13')
               for a, b in (('08:30', '10:30'), ('14:00', '16:00'))])
    return fixed(p, 'tiny-completed', '虚构12人·已完成标准方案')


def golden(p):
    """从手工样本字段展开独立预期，不调用 engine/report。"""
    rows = []
    for e in p['exams']:
        members = [s for s in p['students'] if s['grade'] == e['grade']
                   and (not s['subjects'] or e['subject'] in s['subjects'])]
        positions = p['seats'][e['id']] if e['special_seating'] else p['seats']['base']
        for s in members:
            rows.append(dict(student_id=s['id'], exam_id=e['id'], number=s['number'],
                             room_id=positions[s['id']]['room_id'], seat=positions[s['id']]['seat'],
                             date=e['date'], start=e['start'], end=e['end']))
    return dict(basis='手工固定位置与日程；不使用产品求解器或报表生成预期',
                students=len(p['students']), exam_rows=len(rows),
                base_positions=deepcopy(p['seats']['base']),
                invigilation=deepcopy(p['invigilation']), exam_student_rows=rows)


def clear_results(p):
    p['seats'] = {'base': {}}
    p['invigilation'] = {}
    p['number_sources'] = {}
    p['number_history'] = []
    p['history'] = []
    return p


def table_rows(p, kind):
    rows = []
    for entity in p[kind]:
        out = []
        for key in FIELDS[kind]:
            v = entity.get(key, '')
            if key == 'order': v = '蛇形' if v == 'snake' else '顺排'
            elif key == 'direction': v = '竖排' if v == 'vertical' else '横排'
            elif isinstance(v, bool): v = '是' if v else '否'
            elif isinstance(v, list):
                if key in ('availability', 'unavailable'):
                    v = ';'.join(f"{x['date']} {x['start']}-{x['end']}" for x in v)
                else: v = '、'.join(map(str, v))
            out.append('' if v is None else v)
        rows.append(out)
    return rows


def save_table(path, headers, rows, encoding='utf-8-sig', leading_title=False):
    if path.suffix == '.csv':
        with path.open('w', encoding=encoding, newline='') as f:
            w = csv.writer(f); w.writerow(headers); w.writerows(rows)
        return
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '总名单' if '姓名' in headers else '填写资料'
    if leading_title:
        note = wb.create_sheet('说明', 0); note.append(['本文件只含虚构学生，目标名单在总名单。'])
        ws.append(['验测中学·虚构名单']); ws.append(['第三行是表头。'])
    ws.append(headers); ws.append(rows[0]) if rows else None
    for row in rows[1:]: ws.append(row)
    hr = 3 if leading_title else 1
    ws.freeze_panes = f'A{hr+1}'
    for c in ws[hr]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='334C83')
    for row in ws:
        for c in row:
            if isinstance(c.value, str): c.data_type = 's'; c.number_format = '@'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 22
    wb.save(path); wb.close()


def build_cases():
    cases, plans = [], {}
    base = tiny()

    def add(cid, name, p, steps, codes=(), warnings=(), clean=False, golden_file=False, **expected):
        p = fixed(p, cid.lower().replace('_', '-'), '虚构·'+name)
        plans[cid] = p
        e = dict(required_error_codes=list(codes), required_warning_codes=list(warnings),
                 zero_hard_errors=clean, **expected)
        cases.append(dict(id=cid, name=name, type='plan', path=f'fixtures/{cid}/方案.json',
            counts={k:len(p[k]) for k in KINDS}, steps=steps, expected=e, golden=golden_file,
            execution_status='未运行'))
        return p

    add('C01_tiny_completed', '12人已完成标准方案', base,
        ['打开方案', '检查', '正式导出全部材料', '关闭后重开并比对标准答案'],
        clean=True, golden_file=True, exam_rows=36, duties=6)
    normal = fixed(demo_plan(), 'normal-144', '虚构144人常规方案')
    for s in normal['students']:
        s['name'] = '虚构'+s['name']; s['student_no'] = '00'+s['number']
        s['original'] = {'原始总分': s['score'], '虚构备注': '仅供测试'}
    for t in normal['teachers']: t['name'] = '虚构'+t['name']
    add('C02_normal_144', '144人正常完整流程', normal,
        ['导入四类表', '排座', '排监考', '检查确认', '导出全部材料', '重开'],
        after_solve_zero_errors=True, exam_rows_after_solve=432)
    p = deepcopy(normal)
    for e in p['exams']: e.update(date='', start='', end='')
    add('C03_unscheduled', '未安排日程', p, ['自动日程', '排座与监考', '检查并导出'],
        codes=['exam_time'], after_solve_zero_errors=True)
    p = deepcopy(base); p['exams'][0]['locked'] = True
    p['seats']['base']['s7-001']['locked'] = True
    p['invigilation']['r01|2026-11-12|08:30|10:00'][0]['locked'] = True
    add('C04_locked', '锁定与重排', p, ['记录锁定', '重新排时间、座位和监考', '检查锁定保留', '保存重开'],
        clean=True, preserved=['e7-1', 's7-001', 't01'])
    p = deepcopy(base); p['students'][0]['subjects'] = ['语文', '英语']
    e = p['exams'][1]; e['special_seating'] = True
    mapping = {s['id']:deepcopy(p['seats']['base'][s['id']]) for s in p['students'][1:6]}
    mapping['s7-002'], mapping['s7-003'] = mapping['s7-003'], mapping['s7-002']
    p['seats'][e['id']] = mapping
    add('C05_special_subject', '单科座位与参试子集', p,
        ['检查数学只含5名七年级学生', '比对固定考号', '导出数学材料'], clean=True, golden_file=True,
        exam_rows=35, excluded_math_student='s7-001')
    p = deepcopy(base); p['students'][1]['number'] = p['students'][0]['number']
    add('C06_duplicate_number', '重复考号', p, ['检查定位两位学生', '正式导出拒绝', '改号后复核'], codes=['number_duplicate'])
    p = deepcopy(base); p['students'][0].update(name='', grade='', class_name='')
    add('C07_missing_required', '学生必填资料缺失', p, ['检查', '修复必填字段', '复核'], codes=['student_data'])
    p = deepcopy(base); p['students'][1]['name'] = p['students'][0]['name']; p['students'][0]['score'] = None
    add('C08_same_name_score', '同名与缺成绩', p, ['同名学生均保留', '成绩排序缺成绩最后', '提示不阻止正式导出'],
        clean=True, warnings=['same_name', 'missing_score'])
    p = deepcopy(base); p['rooms'][0]['disabled'] = [6]
    add('C09_disabled_capacity', '停用座位与容量不符', p, ['检查', '减少容量并重新安排', '复核'],
        codes=['room_capacity', 'seat_invalid'])
    p = deepcopy(base); p['seats']['base']['s7-002']['seat'] = 1
    add('C10_duplicate_seat', '同时重复占座', p, ['检查定位学生', '正式导出拒绝', '恢复位置再复核'], codes=['seat_duplicate'])
    p = deepcopy(base); p['exams'][1].update(start='08:30', end='10:00')
    add('C11_exam_overlap', '同年级考试时间重叠', p, ['检查时间冲突', '修改日程后复核'], codes=['grade_overlap', 'student_overlap'])
    p = deepcopy(base); p['rooms'][0]['availability'] = [dict(date='2026-11-12', start='14:00', end='16:00')]
    add('C12_room_unavailable', '考场时段不可用', p, ['检查', '修复时段或换场'], codes=['room_unavailable'])
    p = deepcopy(base); p['teachers'][0]['unavailable'] = [dict(date='2026-11-12', start='08:00', end='12:00')]
    add('C13_teacher_unavailable', '教师时段不可用', p, ['检查', '换合格教师后复核'], codes=['teacher_unavailable'])
    p = deepcopy(base); p['teachers'][0]['subjects'] = ['语文']
    add('C14_subject_avoidance', '任教学科回避', p, ['检查', '人工指定同科教师应拒绝', '换教师'], codes=['subject_avoidance'])
    p = deepcopy(base); p['teachers'][0]['classes'] = ['七1']
    add('C15_class_avoidance', '任教班级回避', p, ['检查', '人工指定本班教师应拒绝', '换教师'], codes=['class_avoidance'])
    p = deepcopy(base); p['invigilation']['r02|2026-11-12|08:30|10:00'][0]['teacher_id'] = 't01'
    add('C16_teacher_overlap', '教师同时两场', p, ['检查', '拒绝新的冲突调整', '改回合格教师'], codes=['teacher_overlap'])
    p = deepcopy(base); p['teachers'] = p['teachers'][:3]; p['invigilation'] = {}
    add('C17_teacher_shortage', '教师不足', p, ['自动监考', '核对partial及具体空缺', '补教师后补排'],
        codes=['invigilator_missing'], solver_stage='invigilators', solver_status=['partial'], missing_per_slot=1)
    p = deepcopy(base); p['seats']['base']['s7-001'], p['seats']['base']['s7-002'] = p['seats']['base']['s7-002'], p['seats']['base']['s7-001']
    add('C18_number_mismatch', '换座后考号不符', p, ['保留旧考号', '预览覆盖生成001前缀', '确认后复核'],
        codes=['number_seat_mismatch'], mismatch_students=2)
    p = deepcopy(base)
    for r in p['rooms']: r.pop('direction')
    p['settings'].pop('number_rule'); p.pop('number_sources'); p.pop('number_history')
    add('C19_legacy', '旧方案兼容', p, ['导入旧方案', '确认默认横排', '打开考号预览', '保存重开'], clean=True)
    # 存储真正缺少新字段的文件，不让 fixed() 的规范化掩盖兼容场景。
    for r in plans['C19_legacy']['rooms']: r.pop('direction', None)
    plans['C19_legacy']['settings'].pop('number_rule', None)
    plans['C19_legacy'].pop('number_sources', None); plans['C19_legacy'].pop('number_history', None)
    large = fixed(demo_plan(large=True), 'large-1500', '虚构1500人规模方案')
    for s in large['students']: s['name'] = '虚构'+s['name']; s['student_no'] = '00'+s['number']
    for t in large['teachers']: t['name'] = '虚构'+t['name']
    large['settings']['mix_grades'] = True
    add('C20_large_mixed', '1500人允许混排', large, ['排座', '排监考', '记录耗时内存', '独立核对覆盖与回避'],
        after_solve_zero_errors=True, exam_rows_after_solve=4500, scale=True)
    p = deepcopy(large); p['settings']['mix_grades'] = False
    add('C21_large_separate_shortage', '1500人禁混且少一场', p, ['排座', '解释至少51场', '不得自动放宽混排'],
        solver_stage='seats', solver_status=['infeasible'], minimum_rooms=51, scale=True)
    p = deepcopy(normal); p['settings'].update(mix_grades=True, mix_groups=[['七年级', '八年级']])
    add('C22_selected_mixing', '七八混排九独立', p, ['排座', '检查七八交插', '九年级不与七八共场', '保存重开'],
        allowed_mixing=[['七年级', '八年级']], after_solve_zero_errors=True)
    p = deepcopy(base); p['rooms'][0].update(direction='vertical', order='snake'); p['rooms'][1].update(order='snake')
    add('C23_layout_routes', '横竖与蛇形布局', p, ['对照物理坐标', '切换双人桌与过道', '停用后重排', '屏幕与打印比对'], clean=True)
    p = deepcopy(base)
    p['seats']['base']['s7-001']['locked'] = True
    p['seats']['base']['s7-002'] = dict(room_id='r01', seat=1, locked=True)
    add('C24_lock_conflict', '锁定位置互相冲突', p, ['重新排座', '明确报告问题', '不得解除锁定绕过'],
        codes=['seat_duplicate'], solver_stage='seats', solver_status=['infeasible', 'invalid'])
    p = deepcopy(base); p['stale'] = ['seats', 'invigilation']
    add('C25_stale', '修改后待复核', p, ['正式导出拒绝', '检查无其他硬问题后确认', '保存重开'], codes=['stale'])
    p = deepcopy(base); p['students'][0]['name'] = '=1+1'; p['students'][1]['name'] = '<script>window.TEST=1</script>'
    add('C26_safe_text', '特殊文本与导出安全', p, ['导出Excel字符串不是公式', '网页按文本显示', '不执行脚本'], clean=True)
    p = deepcopy(base)
    for s in p['students'][6:]:
        pos = p['seats']['base'][s['id']]; pos['room_id'] = 'r01'
        s['number'] = '00801'+str(pos['seat']).zfill(2)
        p['number_sources'][s['id']] = dict(prefix='008', room_id='r01', number=s['number'])
    p['invigilation'] = {}
    for e in p['exams']:
        if e['grade'] == '八年级': e['date'] = '2026-11-14' if e['date'] == '2026-11-12' else '2026-11-15'
        offset = 0 if e['grade'] == '七年级' else 2
        p['invigilation'][f"r01|{e['date']}|{e['start']}|{e['end']}"] = [
            dict(teacher_id=f't{offset+k+1:02}', locked=False) for k in range(2)]
    p['settings']['slots'].extend(dict(date=d, start=a, end=b)
        for d in ('2026-11-14','2026-11-15') for a,b in (('08:30','10:30'),('14:00','16:00')))
    # 1.2.0 起只有流水号方式允许时间不重叠的学生共用固定座位；按考场座位方式每人独占座位。
    p['settings']['number_rule'] = 'serial'
    add('C27_nonoverlap_reuse', '非重叠时段复用座位', p,
        ['流水号方式下检查不同日期同位置合法', '分别按001和008前缀保留号码', '导出材料核对日期'], clean=True, golden_file=True, exam_rows=36)
    return cases, plans


def build_imports(destination, normal):
    path = destination/'fixtures'/'imports'; path.mkdir(parents=True, exist_ok=True)
    cases = []

    def add(ident, name, filename, expected, steps):
        cases.append(dict(id=ident, name=name, type='import', path='fixtures/imports/'+filename,
                          expected=expected, steps=steps, execution_status='未运行'))

    headers, rows = list(FIELDS['students'].values()), table_rows(normal, 'students')
    save_table(path/'多工作表标题行.xlsx', headers, rows, leading_title=True)
    add('I01', '多工作表标题行', '多工作表标题行.xlsx', dict(kind='students', count=144, sheet='总名单', header_row=3), ['预检自动选择总名单和第三行', '确认导入'])
    save_table(path/'学生GB18030.csv', headers, rows, encoding='gb18030')
    add('I02', '中文CSV编码', '学生GB18030.csv', dict(kind='students', count=144, leading_zero=True), ['预检', '导入', '检查中文和前导零'])
    alias_headers = ['学生姓名', '年级', '班级（匀班）', '学籍号', '准考证号', '总分', '参试科目', '第二次总分']
    save_table(path/'别名与原始成绩列.xlsx', alias_headers, [r+[500] for r in rows])
    add('I03', '别名列和额外成绩列', '别名与原始成绩列.xlsx', dict(kind='students', count=144, original_column='第二次总分'), ['预检映射', '导入后原始列保留', '切换成绩排序列'])
    bad = deepcopy(rows[:2]); bad[1][4] = bad[0][4]
    save_table(path/'重复考号.csv', headers, bad)
    add('I04', '导入重复考号', '重复考号.csv', dict(kind='students', count=2, validation_error='number_duplicate'), ['预检可读', '导入后统一检查定位重复', '修复'])
    bad = deepcopy(rows[:1]); bad[0][0] = ''
    save_table(path/'缺姓名.csv', headers, bad)
    add('I05', '导入缺姓名', '缺姓名.csv', dict(kind='students', count=1, validation_error='student_data'), ['预检可读', '导入后检查定位必填缺失'])
    (path/'损坏文件.xlsx').write_bytes(b'fictional damaged xlsx - intentionally not a ZIP archive')
    add('I06', '损坏Excel', '损坏文件.xlsx', dict(kind='students', preview_reject=True), ['预检拒绝并显示中文说明', '当前方案不变'])
    (path/'损坏方案.json').write_text('{"version":1,', encoding='utf-8')
    add('I07', '损坏JSON', '损坏方案.json', dict(plan_import_reject=True), ['打开方案拒绝', '当前方案不变'])
    p = deepcopy(normal); p['version'] = 999
    write_json(path/'不支持版本.json', p)
    add('I08', '不支持的方案版本', '不支持版本.json', dict(normalize_reject=True, plan_import_reject=True), ['打开方案拒绝', '当前方案不变'])
    p = deepcopy(normal)
    heads = list(FIELDS['teachers'].values()); bad = table_rows(p, 'teachers')[:1]; bad[0][-1] = '2026-11-12 时间写错'
    save_table(path/'错误教师时段.csv', heads, bad)
    add('I09', '错误时段文本', '错误教师时段.csv', dict(kind='teachers', count=1, import_reject=True), ['预检可读', '确认导入时拒绝', '当前方案不变'])
    return cases


def acceptance_sheet(destination, cases):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '逐场景验收'
    ws.append(['编号', '场景', '执行状态', '平台与版本', '实际结果', '证据路径', '问题与修复', '复测结果'])
    for c in cases: ws.append([c['id'], c['name'], '未运行', '', '', '', '', ''])
    for i in range(1, 7): ws.append([f'FLOW-{i:02}', '完整流程，见测试方案', '未运行'])
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='334C83')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 25
    note = wb.create_sheet('记录说明'); note.append(['状态：通过／失败／环境阻塞／未运行；不把未运行当通过。'])
    note.append(['异常场景正确拒绝和恢复属于通过；Windows和Mac分别记录。'])
    note.column_dimensions['A'].width = 100
    wb.save(destination/'验收记录模板.xlsx'); wb.close()


def generate(destination):
    cases, plans = build_cases()
    for c in cases:
        p = plans[c['id']]; path = destination/Path(c['path']).parent
        write_json(path/'方案.json', p)
        if c['golden']: write_json(path/'golden_expected.json', golden(normalize_plan(p)))
        if c['id'] in ('C01_tiny_completed', 'C02_normal_144', 'C20_large_mixed'):
            for kind in KINDS:
                for suffix in ('.xlsx', '.csv'):
                    save_table(path/(LABELS[kind]+suffix), list(FIELDS[kind].values()), table_rows(p, kind))
    cases += build_imports(destination, plans['C02_normal_144'])
    # 小布局标准坐标手工写出，防止依赖产品 seat_layout 作为唯一裁判。
    write_json(destination/'fixtures'/'layout_expected.json', dict(
        basis='2行×3桌列、单人桌、无停用，坐标为[行,列]，按座号1至6排列',
        horizontal_straight=[[1,1],[1,2],[1,3],[2,1],[2,2],[2,3]],
        horizontal_snake=[[1,1],[1,2],[1,3],[2,3],[2,2],[2,1]],
        vertical_straight=[[1,1],[2,1],[1,2],[2,2],[1,3],[2,3]],
        vertical_snake=[[1,1],[2,1],[2,2],[1,2],[1,3],[2,3]]))
    acceptance_sheet(destination, cases)
    write_json(destination/'cases.json', dict(version=1, name='考场编排·验测台', seed=42,
        fictional=True, source_date='2026-10-06', scope='资料与预期清单；完整验测执行器尚未实现', cases=cases))
    files = sorted((destination/'fixtures').rglob('*'))
    write_json(destination/'inventory.json', dict(files=[dict(path=str(f.relative_to(destination)),
        bytes=f.stat().st_size, sha256=hashlib.sha256(f.read_bytes()).hexdigest()) for f in files if f.is_file()]))
    return cases


def verify(destination):
    """资料预检：不执行浏览器、桌面、全流程求解或用户目录写操作。"""
    from exam_room import engine, importer, report
    manifest = json.loads((destination/'cases.json').read_text(encoding='utf-8'))
    checks = []

    def check(name, operation):
        try:
            operation(); checks.append(dict(name=name, status='通过'))
        except Exception as exc:
            checks.append(dict(name=name, status='失败', error=f'{type(exc).__name__}: {exc}'))

    def ensure(condition, message):
        if not condition: raise ValueError(message)

    def plan_check(c):
        p = normalize_plan(json.loads((destination/c['path']).read_text(encoding='utf-8')))
        e = c['expected']; v = engine.validate(p)
        codes = {i['code'] for i in v['errors']}; warnings = {i['code'] for i in v['warnings']}
        ensure(all(len(p[k]) == c['counts'][k] for k in KINDS), '资料数量不匹配')
        ensure(set(e['required_error_codes']) <= codes, '未出现预期故障：'+str(set(e['required_error_codes'])-codes))
        ensure(set(e['required_warning_codes']) <= warnings, '未出现预期提示')
        if e['zero_hard_errors']: ensure(not v['errors'], '标准方案出现硬问题：'+str(sorted(codes)))
        if e.get('mismatch_students'):
            ensure(sum(i['code']=='number_seat_mismatch' for i in v['errors']) == e['mismatch_students'], '不匹配学生数量不符')
        if c['golden']:
            g = json.loads((destination/Path(c['path']).parent/'golden_expected.json').read_text(encoding='utf-8'))
            actual = [dict(student_id=r['student_id'], exam_id=r['exam_id'], number=r['考号'], room_id=r['room_id'],
                           seat=r['座号'], date=r['日期'], start=r['时间'].split('—')[0], end=r['时间'].split('—')[1])
                      for r in report.all_rows(p)]
            key = lambda r: (r['exam_id'], r['student_id'])
            ensure(sorted(actual, key=key) == sorted(g['exam_student_rows'], key=key), '报表明细与手工标准答案不符')
            ensure(len(actual) == e['exam_rows'], '标准答案明细行数不符')
            ensure(len(engine.duties(p)) == len(g['invigilation']), '监考场次数不符')

    for c in manifest['cases']:
        if c['type'] == 'plan': check(c['id']+' JSON及预期问题', lambda c=c: plan_check(c))

    def import_check(c):
        e = c['expected']; f = destination/c['path']; blob = f.read_bytes()
        if e.get('plan_import_reject'):
            try: normalize_plan(json.loads(blob.decode('utf-8-sig')))
            except (ValueError, TypeError): return
            raise ValueError('错误方案未被拒绝')
        if e.get('preview_reject'):
            try: importer.preview(blob, f.name, e['kind'])
            except ValueError: return
            raise ValueError('损坏文件未被拒绝')
        pv = importer.preview(blob, f.name, e['kind'])
        ensure(pv['count'] == e['count'], '导入预检数量不匹配')
        if 'sheet' in e: ensure(pv['sheet'] == e['sheet'] and pv['header_row'] == e['header_row'], '表头选择不符')
        if e.get('import_reject'):
            try: importer.import_rows(e['kind'], pv['rows'], pv['mapping'])
            except ValueError: return
            raise ValueError('错误时段未被拒绝')
        rows = importer.import_rows(e['kind'], pv['rows'], pv['mapping'])
        ensure(len(rows) == e['count'], '确认导入数量不匹配')
        if e.get('leading_zero'): ensure(rows[0]['number'] == '700001' and rows[0]['student_no'] == '00700001', '前导零丢失')
        if e.get('original_column'): ensure(rows[0]['original'][e['original_column']] == 500, '原始成绩列未保留')
        if e.get('validation_error'):
            p = tiny(); p[e['kind']] = rows; p['seats'] = {'base':{}}; p['number_sources'] = {}; p['invigilation'] = {}
            ensure(e['validation_error'] in {i['code'] for i in engine.validate(p)['errors']}, '缺失预期导入后问题')

    for c in manifest['cases']:
        if c['type'] == 'import': check(c['id']+' 导入资料预检', lambda c=c: import_check(c))

    def table_check(f, kind, original):
        pv = importer.preview(f.read_bytes(), f.name, kind)
        rows = importer.import_rows(kind, pv['rows'], pv['mapping'])
        ensure(len(rows) == len(original[kind]), '四类资料数量不符')
        for got, source in zip(rows, original[kind]):
            for field in FIELDS[kind]:
                ensure(got[field] == source[field], f'{f.name} 字段{field}不符：{got[field]!r} != {source[field]!r}')

    for cid in ('C01_tiny_completed', 'C02_normal_144', 'C20_large_mixed'):
        path = destination/'fixtures'/cid
        p = normalize_plan(json.loads((path/'方案.json').read_text(encoding='utf-8')))
        for kind in KINDS:
            for suffix in ('.xlsx', '.csv'):
                f = path/(LABELS[kind]+suffix)
                check(cid+' '+f.name, lambda f=f, kind=kind, p=p: table_check(f, kind, p))

    def layout_check():
        g = json.loads((destination/'fixtures'/'layout_expected.json').read_text(encoding='utf-8'))
        for direction in ('horizontal', 'vertical'):
            for order in ('straight', 'snake'):
                r = dict(rows=2, cols=3, per_desk=1, capacity=6, disabled=[], direction=direction, order=order)
                ensure([[x['row'], x['col']] for x in engine.seat_layout(r)] == g[direction+'_'+order], '布局坐标与手工预期不符')

    check('四种布局手工坐标', layout_check)

    def inventory_check():
        inv = json.loads((destination/'inventory.json').read_text(encoding='utf-8'))
        for x in inv['files']:
            blob = (destination/x['path']).read_bytes()
            ensure(len(blob)==x['bytes'] and hashlib.sha256(blob).hexdigest()==x['sha256'], '文件指纹不符：'+x['path'])

    check('文件完整性', inventory_check)
    result = dict(scope='仅测试资料准备检查；不代表软件完整验收或跨平台通过',
                  python=sys.version.split()[0], cases=len(manifest['cases']),
                  checks=len(checks), passed=sum(x['status']=='通过' for x in checks),
                  failed=sum(x['status']=='失败' for x in checks), details=checks)
    out = ROOT/'output'/'testbench-preparation.json'; write_json(out, result)
    print(json.dumps({k:v for k,v in result.items() if k!='details'}, ensure_ascii=False))
    for x in checks:
        if x['status'] == '失败': print(x['name']+': '+x['error'])
    print('准备检查报告：'+str(out))
    return 1 if result['failed'] else 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--verify', action='store_true', help='生成后检查资料，不运行完整验收')
    parser.add_argument('--verify-only', action='store_true', help='只检查现有资料，不重新生成覆盖')
    args = parser.parse_args()
    destination = ROOT/'testkit'
    if not args.verify_only:
        cases = generate(destination)
        print(f'已生成 {len(cases)} 个虚构场景：{destination}')
    return verify(destination) if args.verify or args.verify_only else 0


if __name__ == '__main__':
    raise SystemExit(main())
