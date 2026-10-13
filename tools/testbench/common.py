"""验测记录、隔离服务、独立规则和实际 Excel 核对。"""
from collections import Counter, defaultdict
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime
import io
import json
from pathlib import Path
import time
import traceback
import urllib.request
import zipfile
import re
import unicodedata

ROOT = Path(__file__).resolve().parents[2]
KIT = ROOT/'testkit'


class Blocked(RuntimeError):
    """缺少测试运行环境；不等于通过。"""


def dump(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def scrub_trace(path, tokens):
    """保留浏览器证据，但去掉该次隔离服务的临时口令。"""
    path=Path(path); temporary=path.with_suffix('.redacted.zip')
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as cleaned:
        for item in original.infolist():
            blob=original.read(item.filename)
            for token in tokens: blob=blob.replace(token.encode(),b'FICTIONAL-TOKEN-REDACTED')
            cleaned.writestr(item,blob)
    temporary.replace(path)


class Recorder:
    def __init__(self, root):
        self.root = Path(root).resolve(); self.root.mkdir(parents=True, exist_ok=True); self.cases = []

    def run(self, ident, name, fn):
        directory = self.root/ident; directory.mkdir(parents=True, exist_ok=True)
        started = time.perf_counter()
        result = dict(id=ident, name=name, status='通过', details={})
        try:
            result['details'] = fn() or {}
        except Exception as exc:
            result['status'] = '环境阻塞' if isinstance(exc, Blocked) else '失败'
            result['details'] = dict(error=f'{type(exc).__name__}: {exc}', evidence=str(directory/'error.txt'))
            (directory/'error.txt').write_text(traceback.format_exc(), encoding='utf-8')
        result['seconds'] = round(time.perf_counter()-started, 3)
        self.cases.append(result); dump(self.root/'results.json', self.cases)
        print(f"{result['status']} {ident} {name} ({result['seconds']}秒)", flush=True)
        return result


class Client:
    def __init__(self, url, app): self.url, self.app = url, app

    def raw(self, name, body=None, token=None):
        request = urllib.request.Request(self.url+'api/'+name, data=json.dumps(body or {}).encode(),
            headers={'Content-Type':'application/json', 'X-Token':self.app.token if token is None else token})
        with urllib.request.urlopen(request, timeout=100) as response:
            return json.loads(response.read())

    def call(self, name, body=None):
        result = self.raw(name, body)
        assert result['ok'], f'{name}: {result.get("error")}'
        return result['data']

    def import_plan(self, plan):
        import base64
        return self.call('plan_import', {'data':base64.b64encode(json.dumps(plan).encode()).decode()})

    def download(self, info):
        with urllib.request.urlopen(self.url.rstrip('/')+info['url'], timeout=60) as r: return r.read()

    def solve(self, kind, limit=10, **options):
        self.call('solve', dict(kind=kind, time_limit=limit, **options))
        end = time.monotonic()+120
        while time.monotonic() < end:
            state = self.call('task')
            if state['task']['state'] != 'running': return state
            time.sleep(.08)
        raise AssertionError('后台任务超过120秒监督时限')


@contextmanager
def service(directory, plan=None, license=None):
    from exam_room.server import App, start
    from exam_room.licensing import License
    directory = Path(directory).resolve()
    app = App(root=directory, license=license or License(directory/'license', None, machine='FICTIONAL-TEST'))
    httpd, url, app = start(app=app)
    client = Client(url, app)
    try:
        if plan is not None: client.import_plan(plan)
        yield client
    finally:
        app.cancel.set()
        end = time.monotonic()+10
        while app.task and app.task['state']=='running' and time.monotonic()<end: time.sleep(.05)
        httpd.shutdown(); httpd.server_close()


def load(ident):
    return json.loads((KIT/'fixtures'/ident/'方案.json').read_text(encoding='utf-8'))


def assert_pdf_text(text, p):
    # PDF字体的ToUnicode可能使用康熙部首；姓名也可能合法换行。
    normalized=re.sub(r'\s+','',unicodedata.normalize('NFKC',text))
    for s in p['students']:
        assert unicodedata.normalize('NFKC',s['number']) in normalized,'PDF漏失学生考号：'+s['number']
        assert unicodedata.normalize('NFKC',s['name']) in normalized,'PDF漏失学生姓名：'+s['name']


def event(row):
    datetime.strptime(row['date'], '%Y-%m-%d')
    def minute(s):
        t = datetime.strptime(s, '%H:%M'); return t.hour*60+t.minute
    start, end = minute(row['start']), minute(row['end'])
    assert end > start, '考试起止时间不合法'
    return row['date'], start, end


def overlaps(a, b): return a[0]==b[0] and a[1]<b[2] and b[1]<a[2]


def allows_mixing(p, a, b):
    if a == b: return True
    s = p['settings']; groups = s.get('mix_groups', [])
    return bool(s.get('mix_grades')) and (any(a in g and b in g for g in groups) if groups else True)


def usable(room):
    total = room['rows']*room['cols']*room['per_desk']
    return [n for n in range(1,total+1) if n not in room['disabled']][:room['capacity']]


def expected_rows(p):
    """独立展开实际参试和安排，不调用产品 participants/duties/report。"""
    rooms = {r['id']:r for r in p['rooms']}
    records, duties = [], {}
    for e in p['exams']:
        positions = p['seats'][e['id']] if e.get('special_seating') else p['seats']['base']
        for s in p['students']:
            if s['grade'] != e['grade'] or (s['subjects'] and e['subject'] not in s['subjects']): continue
            assert s['id'] in positions, '漏排学生：'+s['id']
            pos = positions[s['id']]; rid = pos['room_id']
            assert rid in rooms, '未知考场'
            r = rooms[rid]; ident = '|'.join((rid,e['date'],e['start'],e['end']))
            required = e.get('invigilators')
            if required is None: required = r.get('invigilators')
            if required is None: required = p['settings']['invigilators']
            d = duties.setdefault(ident, dict(room_id=rid, date=e['date'], start=e['start'], end=e['end'],
                required=0, subjects=set(), classes=set(), students=set(), exam_ids=set()))
            d['required'] = max(d['required'], required); d['subjects'].add(e['subject'])
            d['classes'].add(s['class_name']); d['students'].add(s['id']); d['exam_ids'].add(e['id'])
            records.append(dict(student_id=s['id'], exam_id=e['id'], room_id=rid, seat=pos['seat'],
                grade=s['grade'], class_name=s['class_name'], name=s['name'], number=s['number'],
                subject=e['subject'], room_name=r['name'], location=r['location'],
                date=e['date'], start=e['start'], end=e['end']))
    return records, duties


def assert_plan(p, allow_shortage=False):
    records, duties = expected_rows(p)
    rooms = {r['id']:r for r in p['rooms']}; teachers = {t['id']:t for t in p['teachers']}
    numbers = [s['number'] for s in p['students']]
    assert all(numbers) and len(set(numbers))==len(numbers), '考号缺失或重复'
    positions, student_times, room_events, teacher_times = (defaultdict(list) for _ in range(4))
    exams = {e['id']:e for e in p['exams']}
    for e in p['exams']: assert event(e)[2]-event(e)[1]==e['duration'], '时长不符'
    for x in records:
        e = event(x); room = rooms[x['room_id']]
        assert x['seat'] in usable(room), '停用或无效座位'
        for previous in positions[x['room_id'],x['seat']]:
            assert not overlaps(event(previous),e), '重复占座'
        positions[x['room_id'],x['seat']].append(x)
        for previous in student_times[x['student_id']]: assert not overlaps(previous,e), '学生考试时间冲突'
        student_times[x['student_id']].append(e)
        if room['availability']:
            assert any(event(w)[0]==e[0] and event(w)[1]<=e[1] and event(w)[2]>=e[2]
                       for w in room['availability']), '考场不在可用时段'
        room_events[x['room_id']].append((x['grade'],e))
    for entries in room_events.values():
        unique = list(set(entries))
        for i,(a,ea) in enumerate(unique):
            for b,eb in unique[i+1:]:
                if overlaps(ea,eb): assert ea==eb and allows_mixing(p,a,b), '违规混排或考场时间冲突'
    for ident,d in duties.items():
        ids = [a['teacher_id'] for a in p['invigilation'].get(ident,[]) if a['teacher_id']]
        assert len(ids)==len(set(ids)), '同场教师重复'
        assert len(ids)<=d['required'], '教师人数过多'
        if not allow_shortage: assert len(ids)==d['required'], '监考教师漏排'
        e = event(d)
        for tid in ids:
            assert tid in teachers, '未知教师'; t = teachers[tid]
            assert not any(overlaps(event(w),e) for w in t['unavailable']), '教师不可用'
            if p['settings']['avoid_subject']: assert not set(t['subjects']) & d['subjects'], '未回避任教学科'
            if p['settings']['avoid_class']: assert not set(t['classes']) & d['classes'], '未回避任教班级'
            assert not any(overlaps(previous,e) for previous in teacher_times[tid]), '教师时间冲突'
            teacher_times[tid].append(e)
    for sid,source in p.get('number_sources',{}).items():
        s = next(s for s in p['students'] if s['id']==sid)
        if s['number'] != source['number']: continue
        pos = p['seats']['base'][sid]; room = rooms[pos['room_id']]
        assert s['number']==source['prefix']+room['code'].zfill(2)+str(pos['seat']).zfill(2), '考号与固定座位不符'
    return dict(students=len(p['students']), exam_rows=len(records), duties=len(duties))


def assert_excel(blob, p, sections, layouts=None, scope=None):
    import openpyxl
    from exam_room.report import SECTIONS
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=False)
    assert set(SECTIONS[k] for k in sections) <= set(wb.sheetnames), '缺失材料工作表'
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row), '用户文本成了Excel公式'
    rows, duties = expected_rows(p)
    scope=scope or {}
    rows=[r for r in rows if (not scope.get('exam_ids') or r['exam_id'] in scope['exam_ids'])
          and (not scope.get('room_ids') or r['room_id'] in scope['room_ids'])
          and (not scope.get('grades') or r['grade'] in scope['grades'])]
    student_cols=['年级','班级','姓名','考号','日期','时间','科目','考场','位置','座号']
    students=[dict(zip(student_cols,(r['grade'],r['class_name'],r['name'],r['number'],r['date'],
        r['start']+'—'+r['end'],r['subject'],r['room_name'],r['location'],str(r['seat']).zfill(2)))) for r in rows]
    teachers={t['id']:t['name'] for t in p['teachers']}; rooms={r['id']:r for r in p['rooms']}
    duty_rows=[]; personal=[]; loads={t['id']:[0,0] for t in p['teachers']}
    for ident,d in duties.items():
        ids=[a['teacher_id'] for a in p['invigilation'].get(ident,[]) if a['teacher_id']]
        values=dict(日期=d['date'],时间=d['start']+'—'+d['end'],考场=rooms[d['room_id']]['name'],
            位置=rooms[d['room_id']]['location'],科目='、'.join(sorted(d['subjects'])),
            人数=len(d['students']),应到监考=d['required'],监考教师='、'.join(teachers[t] for t in ids))
        duty_rows.append(values)
        for tid in ids:
            personal.append(dict(教师=teachers[tid],**values))
            loads[tid][0]+=1; loads[tid][1]+=event(d)[2]-event(d)[1]
    expected={k:students for k in ('students','student_details','classes','personal','tickets','rosters','seating','labels','signins')}
    expected.update(rooms=duty_rows,doors=duty_rows,invigilation=duty_rows,teacher_personal=personal,
        schedule=[dict(年级=e['grade'],科目=e['subject'],日期=e['date'],开始=e['start'],结束=e['end'],时长=e['duration']) for e in p['exams']],
        workload=[dict(教师=teachers[tid],场次=counts[0],监考分钟=counts[1]) for tid,counts in loads.items()])
    for key in sections:
        ws=wb[SECTIONS[key]]
        if key=='room_students':
            # 默认并栏表：逐栏读取重复表头，并将每行绑定到上方考场标题。
            actual=[]; fields=(layouts or p.get('report_layouts',{})).get(key,{}).get('field_order') or ['座号','考号','姓名','班级','年级']
            for col in range(1,ws.max_column+1):
                room_name=None; headers=None; timeline=()
                for n in range(3,ws.max_row+1):
                    value=ws.cell(n,col).value
                    if isinstance(value,str) and ' 考场 · ' in value:
                        room_name=value.split(' 考场 · ',1)[1].split(' · ',1)[0]
                    if value==fields[0]:
                        timeline=tuple(sorted(re.findall(r'\d{4}-\d{2}-\d{2} \d{2}:\d{2}—\d{2}:\d{2}',str(ws.cell(n-1,col).value or ''))))
                        headers=[]
                        for c in range(col,ws.max_column+1):
                            h=ws.cell(n,c).value
                            if not h:break
                            headers.append(h)
                        assert headers==fields,'按考场学生总表字段与请求版式不一致'
                        continue
                    if headers and ws.cell(n,col+headers.index('座号')).data_type=='s' and str(ws.cell(n,col+headers.index('座号')).value or '').isdigit():
                        data=dict(zip(headers,[ws.cell(n,col+i).value for i in range(len(headers))]))
                        actual.append((room_name,timeline,*[data.get(k) for k in fields]))
            fixed=[]
            if scope.get('mode')=='exam':
                seen=set()
                for r in rows:
                    identity=(r['student_id'],r['room_id'],r['date'],r['start'],r['end'])
                    if identity in seen:continue
                    seen.add(identity)
                    values=dict(座号=str(r['seat']).zfill(2),考号=r['number'],姓名=r['name'],班级=r['class_name'],年级=r['grade'])
                    fixed.append((r['room_name'],(r['date']+' '+r['start']+'—'+r['end'],),*[values[k] for k in fields]))
            else:
                for student in p['students']:
                    exams=[e for e in p['exams'] if e['grade']==student['grade'] and not e.get('special_seating') and
                        (not student['subjects'] or e['subject'] in student['subjects']) and
                        (not scope.get('exam_ids') or e['id'] in scope['exam_ids'])]
                    if not exams or (scope.get('grades') and student['grade'] not in scope['grades']):continue
                    pos=p['seats']['base'][student['id']]
                    if scope.get('room_ids') and pos['room_id'] not in scope['room_ids']:continue
                    timeline=tuple(sorted({e['date']+' '+e['start']+'—'+e['end'] for e in exams}))
                    values=dict(座号=str(pos['seat']).zfill(2),考号=student['number'],姓名=student['name'],班级=student['class_name'],年级=student['grade'])
                    fixed.append((rooms[pos['room_id']]['name'],timeline,*[values[k] for k in fields]))
            assert Counter(actual)==Counter(fixed),'按考场学生总表明细不一致（并栏漏行或考场错误）'
            continue
        assert key in expected, '未配置独立核对规则：'+key
        grouped_keys=('classes','personal','tickets','labels','rosters','signins','teacher_personal')
        fields=[]
        for c in ws[2]:
            if c.value is None: break
            fields.append(c.value)
        block_headers=[(n,col) for n in range(3,ws.max_row+1) for col in range(1,ws.max_column+1)
            if ws.cell(n,col).value==fields[0] and [ws.cell(n,col+i).value for i in range(len(fields))]==fields]
        if key in grouped_keys and block_headers:
            actual=[]
            for n,col in block_headers:
                title=ws.cell(n-2,col).value
                for row in range(n+1,ws.max_row+1):
                    values=tuple(ws.cell(row,col+i).value or '' for i in range(len(fields)))
                    if not any(values): break
                    actual.append((title,*values))
                    for label in ('考号','座号'):
                        if label in fields: assert ws.cell(row,col+fields.index(label)).data_type=='s','编号未按文本保存'
            wanted=[]; seen=set()
            for index,r in enumerate(expected[key]):
                if key=='classes': title=r['年级']+' '+r['班级']
                elif key=='teacher_personal': title=r['教师']
                elif key in ('rosters','signins'):
                    title=r['考场']+' · '+r['日期']+' '+r['时间']
                    ident=(title,rows[index]['student_id'])
                    if ident in seen: continue
                    seen.add(ident)
                else: title=r['姓名']+' · '+r['班级']+' · 考号 '+r['考号']
                wanted.append((title,*[r.get(h,'') for h in fields]))
            assert Counter(actual)==Counter(wanted),SECTIONS[key]+'分组明细不一致（漏行、标题身份或位置错误）'
            continue
        headers=[c.value for c in ws[2]]
        actual=[tuple(c.value if c.value is not None else '' for c in row) for row in list(ws.rows)[2:]]
        wanted=[tuple(r.get(h,'') for h in headers) for r in expected[key]]
        assert Counter(actual)==Counter(wanted),SECTIONS[key]+'明细不一致（漏行或号码、日期、位置错误）'
        for h in ('考号','座号'):
            if h in headers:
                index=headers.index(h)
                assert all(row[index].data_type=='s' for row in list(ws.rows)[2:]),'编号未按文本保存'
    wb.close()
    return dict(exam_rows=len(rows), sheets=len(sections),independently_checked_sections=list(sections))
